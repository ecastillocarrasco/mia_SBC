"""Renderizado DOCX premium del modelo de documento común de mIA.

Mismo contrato que PDF/XLSX/PPTX/HTML (``mia.reporting.document``): nunca recalcula ninguna cifra,
solo pinta lo que ya decidió el backend, con un layout dedicado por tipo de bloque (tarjetas KPI,
pares etiqueta/valor, severidad en color, gráfico...) — antes este renderizador trataba todo bloque
como una tabla genérica "Concepto/Valor", sin paridad con los demás formatos.

La paleta de marca (NAVY/GOLD/...) se duplica aquí en vez de importarse de ``mia.reporting.render``
porque ese módulo importa ``to_docx`` de este archivo (``RENDERERS``) — importar en sentido inverso
crearía un ciclo. Mismo criterio que ya usa ``render_html.py`` con su propia copia de la paleta.

Gráficos: a diferencia de Excel/PowerPoint, python-docx no tiene un objeto de gráfico nativo
editable, así que se incrustan como imagen PNG de alta resolución generada con matplotlib en modo
headless, usando la API orientada a objetos (``Figure``/``FigureCanvasAgg``, nunca ``pyplot``, que
mantiene estado global compartido entre hilos) porque este renderizador se invoca concurrentemente
desde varios trabajos de informes (``mia.reporting.job_manager``).
"""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Mm, Pt, RGBColor

from mia.assets import resolve_logo
from mia.reporting.document import Block, Document

NAVY, GOLD, LIGHT, LINE, MUTED = "1B2559", "D8A565", "F6F3EE", "E3DDD3", "6B6F80"
PALETTE = ["1B2559", "D8A565", "3A8F8A", "D9776B", "5B8DEF", "8A6BB0", "7FA37A", "C9B8A3"]
SEVERITY_COLOR = {"blocker": "B3261E", "data_quality": "A66A1E", "info": "3A8F8A"}
SEVERITY_LABEL = {"blocker": "BLOQUEANTE", "data_quality": "CALIDAD DE DATOS", "info": "INFO"}
NARRATIVE_LABEL = {
    "fact": "HECHO VERIFICADO", "interpretation": "INTERPRETACIÓN",
    "hypothesis": "HIPÓTESIS", "recommendation": "RECOMENDACIÓN",
}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _cell_text(value: Any) -> str:
    if value is None:
        return "—"
    if isinstance(value, dict | list):
        return str(value)
    return str(value)


# ------------------------------------------------------------------ helpers OOXML de bajo nivel
def _set_cell_shading(cell: Any, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_border(cell: Any, **edges: dict[str, Any]) -> None:
    """``edges``: p. ej. ``top={"sz": 18, "color": "D8A565"}``. ``sz`` en octavos de punto."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attrs in edges.items():
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), str(attrs.get("val", "single")))
        element.set(qn("w:sz"), str(attrs.get("sz", 6)))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), str(attrs.get("color", "auto")))


def _add_field(paragraph: Any, instr: str, *, cached: str = "1") -> list[Any]:
    """Inserta un campo de Word (p. ej. ``PAGE``) con un valor en caché para que se vea algo incluso
    antes de que Word lo recalcule; ``updateFields`` (ver ``_force_field_update``) lo refresca al abrir."""
    runs = []
    for xml_builder in (
        lambda: _field_run(paragraph, "begin"),
        lambda: _instr_run(paragraph, instr),
        lambda: _field_run(paragraph, "separate"),
        lambda: paragraph.add_run(cached),
        lambda: _field_run(paragraph, "end"),
    ):
        runs.append(xml_builder())
    return runs


def _field_run(paragraph: Any, fld_type: str) -> Any:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), fld_type)
    run._r.append(fld_char)
    return run


def _instr_run(paragraph: Any, instr: str) -> Any:
    run = paragraph.add_run()
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = f" {instr} "
    run._r.append(instr_text)
    return run


def _force_field_update(word: Any) -> None:
    """Word recalcula PAGE/NUMPAGES al abrir el documento, sin que el usuario tenga que pulsar F9."""
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    word.settings.element.append(update)


# ------------------------------------------------------------------ gráficos (matplotlib headless)
def _chart_image(spec: dict[str, Any], width_in: float, *, caption: str | None = None) -> bytes:  # noqa: PLR0912, PLR0915
    from matplotlib.backends.backend_agg import FigureCanvasAgg  # noqa: PLC0415
    from matplotlib.figure import Figure  # noqa: PLC0415

    labels = [str(label) for label in spec.get("labels") or []]
    series = [s for s in spec.get("series", []) if s.get("values")]
    hexes = [f"#{c}" for c in PALETTE]
    kind = spec.get("type")
    horizontal = kind == "horizontal_bar" or (kind not in ("donut", "line", "area") and len(labels) > 8)
    height_in = max(2.4, 0.34 * len(labels) + 1.1) if horizontal else 3.1
    fig = Figure(figsize=(width_in, height_in), dpi=220)
    fig.patch.set_facecolor("white")
    canvas = FigureCanvasAgg(fig)
    ax = fig.add_subplot(111)
    ax.set_facecolor("white")
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    for spine in ("left", "bottom"):
        ax.spines[spine].set_color(f"#{LINE}")
    ax.tick_params(colors=f"#{MUTED}", labelsize=8)

    if kind == "donut" and len(series) == 1:
        values = [max(0.0, float(v or 0)) for v in series[0]["values"]]
        total = sum(values) or 1.0
        ax.pie(
            values, labels=None, colors=[hexes[i % len(hexes)] for i in range(len(values))],
            wedgeprops={"width": 0.45, "edgecolor": "white", "linewidth": 1.5}, startangle=90,
        )
        legend_labels = [f"{label}  ({value / total:.0%})" for label, value in zip(labels, values, strict=True)]
        ax.legend(legend_labels, loc="center left", bbox_to_anchor=(1.02, 0.5), frameon=False, fontsize=8)
        for spine in ax.spines.values():
            spine.set_visible(False)
        ax.set_xticks([])
        ax.set_yticks([])
    elif kind in ("line", "area"):
        x = range(len(labels))
        for i, s in enumerate(series):
            color = hexes[i % len(hexes)]
            values = [float(v or 0) for v in s["values"]]
            ax.plot(x, values, color=color, linewidth=2.2, marker="o", markersize=4,
                    label=str(s.get("name", f"Serie {i + 1}")))
            if kind == "area":
                ax.fill_between(x, values, color=color, alpha=0.15)
        ax.set_xticks(list(x))
        rotation, align = (0, "center") if len(labels) <= 10 else (45, "right")
        ax.set_xticklabels(labels, rotation=rotation, ha=align)
        if len(series) > 1:
            ax.legend(frameon=False, fontsize=8, loc="upper left")
        ax.grid(axis="y", color=f"#{LINE}", linewidth=0.6)
    else:
        n = max(1, len(series))
        positions = range(len(labels))
        if horizontal:
            band = 0.8 / n
            for i, s in enumerate(series):
                offsets = [p + (i - (n - 1) / 2) * band for p in positions]
                ax.barh(offsets, [float(v or 0) for v in s["values"]], height=band * 0.92,
                        color=hexes[i % len(hexes)], label=str(s.get("name", f"Serie {i + 1}")))
            ax.set_yticks(list(positions))
            ax.set_yticklabels(labels)
            ax.invert_yaxis()
            ax.grid(axis="x", color=f"#{LINE}", linewidth=0.6)
        else:
            band = 0.8 / n
            for i, s in enumerate(series):
                offsets = [p + (i - (n - 1) / 2) * band for p in positions]
                ax.bar(offsets, [float(v or 0) for v in s["values"]], width=band * 0.92,
                       color=hexes[i % len(hexes)], label=str(s.get("name", f"Serie {i + 1}")))
            ax.set_xticks(list(positions))
            rotation, align = (0, "center") if len(labels) <= 8 else (45, "right")
            ax.set_xticklabels(labels, rotation=rotation, ha=align)
            ax.grid(axis="y", color=f"#{LINE}", linewidth=0.6)
        if n > 1:
            ax.legend(frameon=False, fontsize=8, loc="upper left" if not horizontal else "lower right")

    if caption:
        ax.set_title(caption, fontsize=9, color=f"#{NAVY}", fontweight="bold", loc="left")
    fig.tight_layout(pad=1.1)
    buffer = io.BytesIO()
    canvas.print_png(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ bloques: un layout por tipo
def _flat_evidence(evidence: dict[str, Any]) -> str:
    return ", ".join(f"{k}: {v}" for k, v in evidence.items()) if evidence else "—"


def _style_run(
    run: Any, *, size: float, color: str, bold: bool = False, italic: bool = False, name: str | None = None
) -> None:
    run.font.size = Pt(size)
    run.font.color.rgb = _rgb(color)
    run.font.bold = bold
    run.font.italic = italic
    if name:
        run.font.name = name


def _add_kpis(word: Any, items: list[tuple[str, Any]], content_width: float) -> None:
    if not items:
        return
    per_row = 4 if len(items) > 8 else max(1, min(3, len(items)))
    rows = [items[i : i + per_row] for i in range(0, len(items), per_row)]
    table = word.add_table(rows=len(rows), cols=per_row)
    table.autofit = False
    for column in table.columns:
        column.width = Inches(content_width / per_row)
    for r, row_items in enumerate(rows):
        for c in range(per_row):
            cell = table.cell(r, c)
            cell.width = Inches(content_width / per_row)
            if c >= len(row_items):
                continue
            label, value = row_items[c]
            _set_cell_shading(cell, LIGHT)
            _set_cell_border(cell, top={"sz": 20, "color": GOLD})
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value_p = cell.paragraphs[0]
            value_p.paragraph_format.space_after = Pt(2)
            _style_run(value_p.add_run(str(value)), size=17, color=NAVY, bold=True, name="Aptos Display")
            label_p = cell.add_paragraph()
            _style_run(label_p.add_run(str(label)), size=8.5, color=MUTED)
    word.add_paragraph()


def _add_section_pairs(word: Any, data: dict[str, Any], content_width: float) -> None:
    if not data:
        return
    table = word.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, width in zip(table.columns, (content_width * 0.32, content_width * 0.68), strict=True):
        column.width = Inches(width)
    for key, value in data.items():
        label_cell, value_cell = table.add_row().cells
        label_cell.width, value_cell.width = Inches(content_width * 0.32), Inches(content_width * 0.68)
        value_cell.text = _cell_text(value)
        _style_run(label_cell.paragraphs[0].add_run(str(key)), size=8.5, color=MUTED, bold=True)
        _style_run(value_cell.paragraphs[0].runs[0], size=9.5, color="293447")
        for cell in (label_cell, value_cell):
            _set_cell_border(cell, bottom={"sz": 4, "color": LINE})
    word.add_paragraph()


def _add_panels(word: Any, items: list[dict[str, Any]], *, kind: str) -> None:
    for item in items:
        table = word.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        _set_cell_shading(cell, LIGHT)
        if kind == "warnings":
            severity = item.get("severity", "info")
            color = SEVERITY_COLOR.get(severity, MUTED)
            label = f"{SEVERITY_LABEL.get(severity, str(severity).upper())} · {item.get('code', '')}"
            message = str(item.get("message", ""))
            evidence = None
        else:
            color = GOLD
            label = NARRATIVE_LABEL.get(item.get("kind"), str(item.get("kind", "")).upper())
            message = str(item.get("text", ""))
            evidence_data = item.get("evidence") or {}
            evidence = f"Evidencia: {_flat_evidence(evidence_data)}" if evidence_data else None
        _set_cell_border(cell, left={"sz": 24, "color": color})
        label_p = cell.paragraphs[0]
        _style_run(label_p.add_run(label), size=9, color=color, bold=True)
        message_p = cell.add_paragraph()
        message_p.paragraph_format.space_before = Pt(3)
        _style_run(message_p.add_run(message), size=9.5, color="222222")
        if evidence:
            evidence_p = cell.add_paragraph()
            evidence_p.paragraph_format.space_before = Pt(3)
            _style_run(evidence_p.add_run(evidence), size=8.5, color=MUTED, italic=True)
        word.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_methodology(word: Any, notes: list[str]) -> None:
    for note in notes:
        paragraph = word.add_paragraph(str(note), style="List Number")
        for run in paragraph.runs:
            _style_run(run, size=9.5, color="293447")


def _add_table(word: Any, columns: list[str], rows: list[list[Any]], content_width: float) -> None:
    if not columns:
        return
    table = word.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for cell, name in zip(table.rows[0].cells, columns, strict=True):
        _set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _style_run(cell.paragraphs[0].add_run(str(name)), size=9, color="FFFFFF", bold=True)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, (cell, value) in enumerate(zip(cells, row, strict=False)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                _set_cell_shading(cell, LIGHT)
            run = cell.paragraphs[0].add_run(_cell_text(value))
            if col_index == 0:
                _style_run(run, size=9, color=NAVY, bold=True)
            else:
                _style_run(run, size=9, color="293447")
    word.add_paragraph()


def _annex_notice(block: Block) -> str | None:
    if block.kind != "annex" or not block.data.get("truncated"):
        return None
    shown, total = len(block.data.get("rows", [])), block.data.get("total_rows")
    return (
        f"Mostrando {shown} de {total} filas — el resto no se ha omitido en silencio, "
        "solo no entra en este documento."
    )


def to_docx(doc: Document, generated: datetime, logo: Path | None = None) -> bytes:  # noqa: PLR0912, PLR0915
    """Genera un Word premium: marca Sunset Beach Club en cada página, numeración de páginas real,
    tarjetas KPI, severidad en color y gráficos incrustados como imagen de alta resolución."""
    from mia.reporting.render import _chart_of, _chart_segments, _rows  # noqa: PLC0415

    word = WordDocument()
    section = word.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin, section.bottom_margin = Inches(1.15), Inches(0.75)
    section.header_distance, section.footer_distance = Inches(0.35), Inches(0.35)
    content_width = Emu(section.page_width - section.left_margin - section.right_margin).inches

    styles = word.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = _rgb("2D3447")

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.tab_stops.add_tab_stop(Inches(content_width), WD_TAB_ALIGNMENT.RIGHT)
    logo_path = resolve_logo(logo)
    if logo_path is not None and logo_path.is_file():
        header.add_run().add_picture(str(logo_path), width=Inches(0.8))
    eyebrow = header.add_run("\tSUNSET BEACH CLUB  ·  INFORME")
    _style_run(eyebrow, size=8.5, color=MUTED, bold=True)
    rule_table = section.header.add_table(rows=1, cols=1, width=Inches(content_width))
    rule_cell = rule_table.rows[0].cells[0]
    _set_cell_shading(rule_cell, GOLD)
    rule_table.rows[0].height = Pt(2.2)

    eyebrow_p = word.add_paragraph()
    _style_run(eyebrow_p.add_run("ANÁLISIS CORPORATIVO  ·  INFORME"), size=9, color=GOLD, bold=True)
    title = word.add_heading(doc.title, 0)
    for run in title.runs:
        run.font.name = "Aptos Display"
        run.font.size = Pt(23)
        run.font.bold = True
        run.font.color.rgb = _rgb(NAVY)
    meta = "  ·  ".join(x for x in [doc.subtitle, *doc.badges] if x)
    if meta:
        meta_p = word.add_paragraph()
        _style_run(meta_p.add_run(meta), size=10, color=MUTED)
    gen_p = word.add_paragraph()
    _style_run(gen_p.add_run(f"Generado el {generated:%d/%m/%Y %H:%M}"), size=8, color=MUTED, italic=True)
    word.add_paragraph()

    has_explicit_chart = any(b.kind == "chart" for b in doc.blocks)
    for block in doc.blocks:
        heading = word.add_heading(block.title, level=1)
        for run in heading.runs:
            run.font.name = "Aptos Display"
            run.font.size = Pt(13)
            run.font.color.rgb = _rgb(NAVY)

        if block.kind == "kpis":
            _add_kpis(word, list(block.data.items()), content_width)
            continue
        if block.kind == "section":
            _add_section_pairs(word, block.data, content_width)
            continue
        if block.kind == "warnings":
            _add_panels(word, block.data.get("items", []), kind="warnings")
            continue
        if block.kind == "narrative":
            _add_panels(word, block.data.get("statements", []), kind="narrative")
            continue
        if block.kind == "methodology":
            _add_methodology(word, block.data.get("notes", []))
            continue

        # Bloque "table"/"annex" con un resumen de indicadores adjunto (p. ej. un informe con una
        # entidad por bloque — persona, proyecto…): tarjetas KPI bajo el MISMO título, antes de su
        # propia tabla — nunca un bloque "kpis" aparte (ver sidecar equivalente en to_pdf/to_xlsx).
        if isinstance(block.data, dict) and isinstance(block.data.get("kpis"), dict) and block.data["kpis"]:
            _add_kpis(word, list(block.data["kpis"].items()), content_width)

        if notice := _annex_notice(block):
            notice_p = word.add_paragraph()
            _style_run(notice_p.add_run(notice), size=8.5, color=GOLD, italic=True)

        if block.kind == "chart":
            spec = block.data
        elif not has_explicit_chart:
            spec = _chart_of(block)
        else:
            spec = None
        if spec:
            for segment in _chart_segments(spec):
                caption = segment.get("title") if segment is not spec else None
                image = _chart_image(segment, content_width, caption=caption)
                word.add_picture(io.BytesIO(image), width=Inches(content_width))
            word.add_paragraph()

        if block.kind != "chart":
            columns, rows = _rows(block)
            _add_table(word, columns, rows, content_width)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(footer.add_run("Sunset Beach Club · Documento generado por mIA · Página "), size=8, color=MUTED)
    page_run = _add_field(footer, "PAGE")[-2]
    _style_run(page_run, size=8, color=MUTED, bold=True)
    de_run = footer.add_run(" de ")
    _style_run(de_run, size=8, color=MUTED)
    total_run = _add_field(footer, "NUMPAGES")[-2]
    _style_run(total_run, size=8, color=MUTED, bold=True)
    _force_field_update(word)

    output = io.BytesIO()
    word.save(output)
    return output.getvalue()
