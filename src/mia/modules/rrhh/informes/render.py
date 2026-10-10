"""Plantilla corporativa de informes (PDF, Excel y PowerPoint) a partir del modelo de documento.

Cabecera con el logo de Sunset Beach Club en todas las páginas, filete dorado, secciones con marca, tarjetas de
indicadores y gráficos incrustados cuando el resultado es un recuento por categorías o una serie temporal.
"""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Any

from mia.modules.rrhh.charts import chart_spec
from mia.modules.rrhh.informes.document import Block, Document

NAVY = "1B2559"
GOLD = "D8A565"
LIGHT = "F6F3EE"
LINE = "E3DDD3"
MUTED = "6B6F80"
PALETTE = ["1B2559", "D8A565", "3A8F8A", "D9776B", "5B8DEF", "8A6BB0", "7FA37A", "C9B8A3"]


def _rows(block: Block) -> tuple[list[str], list[list[Any]]]:
    if block.kind == "table":
        return [str(c) for c in block.data.get("columns", [])], list(block.data.get("rows", []))
    return ["Concepto", "Valor"], [[k, v] for k, v in block.data.items()]


def _chart_of(block: Block) -> dict[str, Any] | None:
    if block.kind != "table":
        return None
    columns, rows = _rows(block)
    return chart_spec(block.title, columns, rows)


# ------------------------------------------------------------------ Excel
def to_xlsx(doc: Document, generated: datetime, logo: Path | None = None) -> bytes:  # noqa: PLR0915
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference  # noqa: PLC0415
    from openpyxl.drawing.image import Image  # noqa: PLC0415
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    hair = Side(style="thin", color=LINE)
    wb = Workbook()
    wb.remove(wb.active)
    for index, block in enumerate(doc.blocks or [Block("section", "Informe", {})], 1):
        ws = wb.create_sheet((block.title or f"Hoja {index}")[:28].replace("/", "-") or f"Hoja {index}")
        ws.sheet_view.showGridLines = False
        header_row = 6
        if logo:
            image = Image(str(logo))
            ratio = image.width / max(1, image.height)
            image.height, image.width = 62, 62 * ratio
            ws.add_image(image, "A1")
        ws["C2"] = doc.title
        ws["C2"].font = Font(bold=True, size=16, color=NAVY)
        ws["C3"] = block.title if block.title != doc.title else (doc.subtitle or "")
        ws["C3"].font = Font(size=10, color=MUTED)
        for col in range(1, 13):
            ws.cell(row=5, column=col).border = Border(top=Side(style="medium", color=GOLD))
        columns, rows = _rows(block)
        for c, name in enumerate(columns, 1):
            cell = ws.cell(row=header_row, column=c, value=name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r, row in enumerate(rows, header_row + 1):
            for c, value in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=value)
                cell.border = Border(bottom=hair)
                if (r - header_row) % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=LIGHT)
                if isinstance(value, int | float) and not isinstance(value, bool):
                    cell.number_format = "#,##0.##"
                    cell.alignment = Alignment(horizontal="right")
        ws.column_dimensions["A"].width = 10
        for c, name in enumerate(columns, 1):
            width = max([len(str(name))] + [len(str(row[c - 1] or "")) for row in rows[:500] if len(row) >= c]) + 3
            ws.column_dimensions[get_column_letter(c)].width = max(
                ws.column_dimensions[get_column_letter(c)].width or 0, min(max(width, 10), 60)
            )
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        if rows and columns:
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{header_row + len(rows)}"
        spec = _chart_of(block)
        if spec and rows:
            chart: Any = {"donut": PieChart, "line": LineChart}.get(spec["type"], BarChart)()
            chart.title = spec["title"]
            data = Reference(ws, min_col=2, min_row=header_row, max_row=header_row + len(rows))
            cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + len(rows))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            if spec["type"] == "bar":
                chart.type = "bar"
                chart.legend = None
            chart.height, chart.width = 9, 16
            ws.add_chart(chart, f"{get_column_letter(len(columns) + 2)}{header_row}")
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToHeight = 0
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _smart_col_widths(columns: list[str], rows: list[list[Any]], total: float) -> list[float]:
    """Distribute column widths proportionally: text-heavy first column gets more space,
    short numeric columns stay narrow."""
    import re as _re  # noqa: PLC0415

    n = len(columns)
    if n <= 1:
        return [total]
    sample = rows[:80]
    max_lens = [
        max(len(str(columns[c])), max((len(str(r[c])) for r in sample if len(r) > c), default=0))
        for c in range(n)
    ]
    # Numeric-looking columns cap at a shorter max
    for c in range(1, n):
        is_numeric = all(
            _re.match(r"^[\d\s€%/,.\-—]+$", str(r[c])) for r in sample[:20] if len(r) > c and str(r[c]).strip()
        )
        if is_numeric:
            max_lens[c] = min(max_lens[c], 12)
    total_chars = sum(max_lens) or 1
    raw = [total * length / total_chars for length in max_lens]
    # First column minimum 25% of total, max 55%
    raw[0] = max(raw[0], total * 0.25)
    raw[0] = min(raw[0], total * 0.55)
    # Rescale remaining to fill exactly
    rest = total - raw[0]
    rest_sum = sum(raw[1:]) or 1
    raw[1:] = [v * rest / rest_sum for v in raw[1:]]
    return raw


# ------------------------------------------------------------------ PDF
def _pdf_chart(spec: dict[str, Any], width: float) -> Any:
    from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart  # noqa: PLC0415
    from reportlab.graphics.charts.linecharts import HorizontalLineChart  # noqa: PLC0415
    from reportlab.graphics.charts.piecharts import Pie  # noqa: PLC0415
    from reportlab.graphics.shapes import Drawing, String  # noqa: PLC0415
    from reportlab.lib import colors  # noqa: PLC0415

    labels, values = spec["labels"], spec["series"][0]["values"]
    hexes = [colors.HexColor(f"#{c}") for c in PALETTE]
    if spec["type"] == "donut":
        drawing = Drawing(width, 150)
        pie = Pie()
        pie.x, pie.y, pie.width, pie.height = 20, 15, 120, 120
        pie.data, pie.labels = values, None
        pie.innerRadiusFraction = 0.55
        for i in range(len(values)):
            pie.slices[i].fillColor = hexes[i % len(hexes)]
            pie.slices[i].strokeColor = colors.white
        drawing.add(pie)
        total = sum(values) or 1
        for i, (label, value) in enumerate(zip(labels, values, strict=True)):
            y = 125 - i * 18
            drawing.add(
                String(
                    170,
                    y,
                    f"{label} — {value:g} ({value / total:.0%})",
                    fontName="Helvetica",
                    fontSize=9,
                    fillColor=colors.HexColor(f"#{NAVY}"),
                )
            )
        return drawing
    if spec["type"] == "line":
        drawing = Drawing(width, 170)
        line = HorizontalLineChart()
        line.x, line.y, line.width, line.height = 40, 30, width - 70, 115
        line.data = [values]
        line.categoryAxis.categoryNames = labels
        line.categoryAxis.labels.fontSize = 8
        line.valueAxis.labels.fontSize = 8
        line.lines[0].strokeColor = colors.HexColor(f"#{NAVY}")
        line.lines[0].strokeWidth = 2
        drawing.add(line)
        return drawing
    height = max(120, 16 * len(labels) + 30)
    drawing = Drawing(width, height)
    bars = HorizontalBarChart() if len(labels) > 6 else VerticalBarChart()
    bars.x, bars.y, bars.width, bars.height = 110 if len(labels) > 6 else 40, 20, width - 150, height - 40
    bars.data = [list(reversed(values)) if len(labels) > 6 else values]
    bars.categoryAxis.categoryNames = list(reversed(labels)) if len(labels) > 6 else labels
    bars.categoryAxis.labels.fontSize = 8
    bars.valueAxis.labels.fontSize = 8
    bars.valueAxis.valueMin = 0
    bars.bars[0].fillColor = colors.HexColor(f"#{NAVY}")
    bars.bars.strokeColor = None
    bars.barLabelFormat = "%g"
    bars.barLabels.fontSize = 8
    bars.barLabels.nudge = 8 if len(labels) <= 6 else 10
    bars.barLabels.boxAnchor = "w" if len(labels) > 6 else "s"
    drawing.add(bars)
    return drawing


def to_pdf(doc: Document, generated: datetime, logo: Path | None = None) -> bytes:  # noqa: PLR0915
    from reportlab.lib import colors  # noqa: PLC0415
    from reportlab.lib.pagesizes import A4, landscape  # noqa: PLC0415
    from reportlab.lib.styles import ParagraphStyle  # noqa: PLC0415
    from reportlab.lib.units import cm  # noqa: PLC0415
    from reportlab.lib.utils import ImageReader  # noqa: PLC0415
    from reportlab.platypus import (  # noqa: PLC0415
        KeepTogether,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    hexc = lambda h: colors.HexColor(f"#{h}")  # noqa: E731
    widest = max((len(_rows(b)[0]) for b in doc.blocks if b.kind == "table"), default=2)
    page = landscape(A4) if widest > 4 else A4
    margin = 1.6 * cm
    logo_h = 1.5 * cm
    logo_w = 0.0
    if logo:
        iw, ih = ImageReader(str(logo)).getSize()
        logo_w = logo_h * iw / ih

    def decorate(canvas: Any, pdf_doc: Any) -> None:
        canvas.saveState()
        if logo:
            canvas.drawImage(
                str(logo),
                margin,
                page[1] - margin - logo_h + 0.2 * cm,
                width=logo_w,
                height=logo_h,
                mask="auto",
                preserveAspectRatio=True,
            )
        canvas.setStrokeColor(hexc(GOLD))
        canvas.setLineWidth(1.2)
        y = page[1] - margin - logo_h - 0.05 * cm
        canvas.line(margin, y, page[0] - margin, y)
        canvas.setStrokeColor(hexc(LINE))
        canvas.setLineWidth(0.5)
        canvas.line(margin, margin - 0.25 * cm, page[0] - margin, margin - 0.25 * cm)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(hexc(MUTED))
        canvas.drawRightString(page[0] - margin, margin - 0.65 * cm, f"Página {canvas.getPageNumber()}")
        canvas.restoreState()

    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=page,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin + logo_h + 0.6 * cm,
        bottomMargin=margin + 0.3 * cm,
        title=doc.title,
        author="Sunset Beach Club",
    )
    title_style = ParagraphStyle(
        "t", fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=hexc(NAVY), spaceAfter=4
    )
    sub_style = ParagraphStyle("s", fontName="Helvetica", fontSize=11, leading=14, textColor=hexc(MUTED))
    h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=hexc(NAVY))
    cell = ParagraphStyle("c", fontName="Helvetica", fontSize=8.5, leading=10.5, textColor=hexc("222222"))
    head = ParagraphStyle("h", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    label_style = ParagraphStyle("l", parent=cell, fontName="Helvetica-Bold", textColor=hexc(MUTED), fontSize=8)
    value_style = ParagraphStyle("v", parent=cell, fontSize=9.5, leading=12)
    kpi_value = ParagraphStyle(
        "kv", fontName="Helvetica-Bold", fontSize=18, leading=21, textColor=hexc(NAVY), alignment=1
    )
    kpi_label = ParagraphStyle("kl", fontName="Helvetica", fontSize=8, leading=10, textColor=hexc(MUTED), alignment=1)

    def para(value: Any, style: Any) -> Paragraph:
        text = "" if value is None else str(value)
        return Paragraph(text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), style)

    width = page[0] - 2 * margin

    def heading(text: str) -> Table:
        bar = Table([[para(text, h2)]], colWidths=[width - 0.35 * cm])
        wrapper = Table([["", bar]], colWidths=[0.2 * cm, width - 0.2 * cm])
        wrapper.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (0, 0), hexc(GOLD)),
                    ("LEFTPADDING", (1, 0), (1, 0), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 1),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                ]
            )
        )
        return wrapper

    story: list[Any] = [para(doc.title, title_style)]
    meta = "  ·  ".join(x for x in [doc.subtitle, *doc.badges] if x)
    if meta:
        story.append(para(meta, sub_style))
    story.append(Spacer(1, 0.5 * cm))
    for block in doc.blocks:
        columns, rows = _rows(block)
        if not rows:
            continue
        same_title = block.title == doc.title
        parts: list[Any] = [Spacer(1, 0.05 * cm) if same_title else heading(block.title), Spacer(1, 0.25 * cm)]
        if block.kind == "kpis":
            items = list(block.data.items())
            # 5 per row when there are many KPIs (executive reports); else 4.
            per_row = 5 if len(items) > 8 else min(4, max(1, len(items)))
            kpi_w = (width - 0.3 * cm * per_row) / per_row
            cards = [
                [
                    Table([[para(v, kpi_value)], [para(k, kpi_label)]], colWidths=[kpi_w])
                    for k, v in items[i : i + per_row]
                ]
                for i in range(0, len(items), per_row)
            ]
            cards = [[*card_row, *([""] * (per_row - len(card_row)))] for card_row in cards]
            grid = Table(cards, colWidths=[width / per_row] * per_row)
            grid.setStyle(
                TableStyle(
                    [
                        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ]
                )
            )
            for row_index in range(len(cards)):
                for col_index in range(per_row):
                    if cards[row_index][col_index] != "":
                        grid.setStyle(
                            TableStyle(
                                [
                                    ("BACKGROUND", (col_index, row_index), (col_index, row_index), hexc(LIGHT)),
                                    ("BOX", (col_index, row_index), (col_index, row_index), 0.8, hexc(LINE)),
                                    ("TOPPADDING", (col_index, row_index), (col_index, row_index), 8),
                                    ("BOTTOMPADDING", (col_index, row_index), (col_index, row_index), 8),
                                ]
                            )
                        )
            parts.append(grid)
        elif block.kind == "section":
            pairs = Table(
                [[para(k, label_style), para(v, value_style)] for k, v in block.data.items()],
                colWidths=[width * 0.32, width * 0.68],
            )
            pairs.setStyle(
                TableStyle(
                    [
                        ("LINEBELOW", (0, 0), (-1, -1), 0.4, hexc(LINE)),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ]
                )
            )
            parts.append(pairs)
        else:
            if spec := _chart_of(block):
                parts += [_pdf_chart(spec, width), Spacer(1, 0.2 * cm)]
            col_widths = _smart_col_widths(columns, rows, width)
            first_col_style = ParagraphStyle("fc", parent=cell, fontName="Helvetica-Bold", textColor=hexc(NAVY))

            def _data_row(row: list[Any]) -> list[Any]:
                return [para(v, first_col_style if i == 0 else cell) for i, v in enumerate(row)]

            grid = Table(
                [[para(c, head) for c in columns]] + [_data_row(row) for row in rows],
                colWidths=col_widths,
                repeatRows=1,
            )
            grid.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), hexc(NAVY)),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, hexc(LIGHT)]),
                        ("LINEBELOW", (0, 1), (-1, -1), 0.3, hexc(LINE)),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                        ("LEFTPADDING", (0, 0), (-1, -1), 7),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                        ("LINEAFTER", (0, 0), (0, -1), 0.5, hexc(LINE)),
                    ]
                )
            )
            parts.append(grid)
        story += [KeepTogether(parts[:2] + parts[2:3]), *parts[3:], Spacer(1, 0.55 * cm)]
    pdf.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buffer.getvalue()


# ------------------------------------------------------------------ PowerPoint
def to_pptx(doc: Document, generated: datetime, logo: Path | None = None) -> bytes:  # noqa: PLR0915
    from pptx import Presentation  # noqa: PLC0415
    from pptx.dml.color import RGBColor  # noqa: PLC0415
    from pptx.util import Inches, Pt  # noqa: PLC0415

    navy, gold = RGBColor.from_string(NAVY), RGBColor.from_string(GOLD)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    counter = {"n": 0}

    def header(slide: Any, text: str) -> None:
        counter["n"] += 1
        if logo:
            slide.shapes.add_picture(str(logo), Inches(0.4), Inches(0.15), height=Inches(0.85))
        line = slide.shapes.add_shape(1, Inches(0.4), Inches(1.1), Inches(12.5), Inches(0.03))
        line.fill.solid()
        line.fill.fore_color.rgb = gold
        line.line.fill.background()
        box = slide.shapes.add_textbox(Inches(1.7), Inches(0.3), Inches(11), Inches(0.6)).text_frame
        box.text = text
        run = box.paragraphs[0].runs[0]
        run.font.size, run.font.bold, run.font.color.rgb = Pt(26), True, navy
        foot = slide.shapes.add_textbox(Inches(12.2), Inches(7.05), Inches(0.8), Inches(0.3)).text_frame
        foot.text = str(counter["n"])
        foot.paragraphs[0].runs[0].font.size = Pt(9)
        foot.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    cover = prs.slides.add_slide(blank)
    counter["n"] += 1
    if logo:
        cover.shapes.add_picture(str(logo), Inches(5.65), Inches(0.9), height=Inches(2.1))
    bar = cover.shapes.add_shape(1, Inches(4.9), Inches(3.35), Inches(3.5), Inches(0.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = gold
    bar.line.fill.background()
    title = cover.shapes.add_textbox(Inches(1), Inches(3.6), Inches(11.3), Inches(1.4)).text_frame
    title.word_wrap = True
    title.text = doc.title
    title.paragraphs[0].alignment = 2
    title.paragraphs[0].runs[0].font.size, title.paragraphs[0].runs[0].font.bold = Pt(36), True
    title.paragraphs[0].runs[0].font.color.rgb = navy
    if doc.subtitle:
        sub = cover.shapes.add_textbox(Inches(1), Inches(5.0), Inches(11.3), Inches(0.6)).text_frame
        sub.text = doc.subtitle
        sub.paragraphs[0].alignment = 2
        sub.paragraphs[0].runs[0].font.size = Pt(16)
        sub.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    for block in doc.blocks:
        if block.kind == "kpis":
            _kpi_slide(prs, blank, header, block)
            continue
        if spec := _chart_of(block):
            _chart_slide(prs, blank, header, block.title, spec)
        _table_slides(prs, blank, header, block)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _chart_slide(prs: Any, layout: Any, header: Any, title: str, spec: dict[str, Any]) -> None:
    from pptx.chart.data import CategoryChartData  # noqa: PLC0415
    from pptx.dml.color import RGBColor  # noqa: PLC0415
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION  # noqa: PLC0415
    from pptx.util import Inches  # noqa: PLC0415

    slide = prs.slides.add_slide(layout)
    header(slide, title)
    data = CategoryChartData()
    data.categories = spec["labels"]
    data.add_series(spec["series"][0]["name"], spec["series"][0]["values"])
    kind = {"donut": XL_CHART_TYPE.DOUGHNUT, "line": XL_CHART_TYPE.LINE_MARKERS}.get(
        spec["type"], XL_CHART_TYPE.BAR_CLUSTERED
    )
    chart = slide.shapes.add_chart(kind, Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.4), data).chart
    chart.has_legend = spec["type"] == "donut"
    if chart.has_legend:
        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.RIGHT, False
    plot = chart.plots[0]
    plot.has_data_labels = True
    if spec["type"] != "donut":
        series = plot.series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(NAVY)
        if spec["type"] == "bar":
            chart.category_axis.reverse_order = True


def _kpi_slide(prs: Any, layout: Any, header: Any, block: Block) -> None:
    from pptx.dml.color import RGBColor  # noqa: PLC0415
    from pptx.util import Inches, Pt  # noqa: PLC0415

    navy = RGBColor.from_string(NAVY)
    slide = prs.slides.add_slide(layout)
    header(slide, block.title)
    for i, (label, value) in enumerate(list(block.data.items())[:12]):
        x, y = Inches(0.5 + (i % 4) * 3.15), Inches(1.5 + (i // 4) * 1.8)
        card = slide.shapes.add_shape(1, x, y, Inches(2.95), Inches(1.6))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor.from_string(LIGHT)
        card.line.color.rgb = RGBColor.from_string(LINE)
        tf = card.text_frame
        tf.text = str(value)[:40]
        tf.paragraphs[0].runs[0].font.size = Pt(26 if len(str(value)) < 12 else 13)
        tf.paragraphs[0].runs[0].font.bold = True
        tf.paragraphs[0].runs[0].font.color.rgb = navy
        tf.add_paragraph().text = str(label)
        tf.paragraphs[1].runs[0].font.size = Pt(11)
        tf.paragraphs[1].runs[0].font.color.rgb = RGBColor.from_string(MUTED)


def _table_slides(prs: Any, layout: Any, header: Any, block: Block) -> None:
    from pptx.dml.color import RGBColor  # noqa: PLC0415
    from pptx.util import Inches, Pt  # noqa: PLC0415

    columns, rows = _rows(block)
    for start in range(0, len(rows), 12):
        chunk = rows[start : start + 12]
        slide = prs.slides.add_slide(layout)
        header(slide, block.title if start == 0 else f"{block.title} (cont.)")
        shape = slide.shapes.add_table(
            len(chunk) + 1, len(columns), Inches(0.4), Inches(1.4), Inches(12.5), Inches(0.4) * (len(chunk) + 1)
        )
        grid = shape.table
        for c, name in enumerate(columns):
            grid.cell(0, c).text = str(name)
            grid.cell(0, c).fill.solid()
            grid.cell(0, c).fill.fore_color.rgb = RGBColor.from_string(NAVY)
        for r, row in enumerate(chunk, 1):
            for c in range(len(columns)):
                cell = grid.cell(r, c)
                cell.text = "" if c >= len(row) or row[c] is None else str(row[c])[:120]
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(LIGHT if r % 2 == 0 else "FFFFFF")
        for r in range(len(chunk) + 1):
            for c in range(len(columns)):
                for paragraph in grid.cell(r, c).text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(11 if r == 0 else 10)
                        run.font.color.rgb = RGBColor(255, 255, 255) if r == 0 else RGBColor(0x22, 0x22, 0x22)
                        run.font.bold = r == 0


RENDERERS = {"pdf": to_pdf, "xlsx": to_xlsx, "pptx": to_pptx}
