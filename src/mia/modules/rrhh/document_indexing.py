"""Indexación de documentos sensibles: descarga, extracción de texto, fragmentos y embeddings.

Utilidades genéricas (extracción, OCR, fragmentado, empaquetado de vectores) compartidas por todos
los dominios que indexan documentos — CVs (``curriculums/indexer.py``) y, desde la decisión explícita
de anular NOM-003/FIR-001 como "nunca se guarda", también nóminas y documentos firmados
(``DocumentIndexer``, generalizado por dominio sobre ``HrDocumentText``/``HrDocumentChunk``).

Se ejecuta tras la sincronización. Solo procesa documentos nuevos o modificados (por lotes), así que
el coste de embeddings se paga una vez por documento.
"""

from __future__ import annotations

import asyncio
import io
import logging
import re
import shutil
from functools import lru_cache
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any

import numpy as np
from sqlalchemy import and_, delete, or_, select

from mia.connectors.sharepoint import SharePointClient, SharePointError
from mia.data import Database, HrDocumentChunk, HrDocumentText
from mia.llm import LLMProvider
from mia.modules.rrhh.payroll_embedding_privacy import (
    PAYROLL_EMBEDDING_POLICY_VERSION,
    project_payroll_embedding,
)
from mia.observability import get_logger

log = get_logger(__name__)
logging.getLogger("pypdf").setLevel(logging.ERROR)  # PDF mal formados: el estado ya lo refleja

CHUNK_CHARS = 1200
CHUNK_OVERLAP = 150
EMBED_BATCH = 64


@dataclass(slots=True)
class IndexReport:
    processed: int = 0
    indexed: int = 0
    failed: int = 0
    removed: int = 0
    pending: int = 0
    statuses: dict[str, int] = field(default_factory=dict)


_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")
_IMAGE_MAGIC = (b"\xff\xd8\xff", b"\x89PNG", b"GIF8", b"II*\x00", b"MM\x00*")


OCR_INSTRUCTION = (
    "Transcribe literalmente el texto de este documento (todos los datos y conceptos, tal como aparecen). "
    "No añadas comentarios ni valoraciones."
)
_MEDIA = {b"\xff\xd8\xff": "image/jpeg", b"\x89PNG": "image/png", b"GIF8": "image/gif"}


def image_media_type(content: bytes) -> str | None:
    return next((media for magic, media in _MEDIA.items() if content.startswith(magic)), None)


OCR_MAX_PAGES = 3


def render_pdf_pages(content: bytes, max_pages: int) -> list[bytes]:
    """Primeras páginas de un PDF como PNG (para documentos escaneados)."""
    import pypdfium2 as pdfium  # noqa: PLC0415

    pdf = pdfium.PdfDocument(content)
    images = []
    try:
        for index in range(min(len(pdf), max_pages)):
            bitmap = pdf[index].render(scale=1.6)
            buffer = io.BytesIO()
            bitmap.to_pil().save(buffer, format="PNG")
            images.append(buffer.getvalue())
    finally:
        pdf.close()
    return images


def clean_text(text: str) -> str:
    """Quita caracteres de control (p. ej. NUL), que PostgreSQL no admite en texto."""
    return " ".join(_CONTROL.sub(" ", text).split())


# OCR exclusivo de la indexacion documental. No se usa un LLM para leer imagenes.
# El limite es un control de seguridad: si no cabe todo el PDF, NO se indexa parcialmente.
LOCAL_OCR_MAX_SCANNED_PAGES = 10
LOCAL_OCR_TIMEOUT_SECONDS = 45
LOCAL_OCR_MAX_PIXELS = 25_000_000
LOCAL_OCR_MIN_NATIVE_CHARS = 30


class LocalOCRError(RuntimeError):
    """No se puede completar el OCR sin recurrir a proveedores externos."""


@lru_cache(maxsize=1)
def _local_tesseract():
    """Localiza Tesseract y comprueba el modelo espanol una vez por proceso."""
    try:
        import pytesseract
    except ImportError as exc:
        raise LocalOCRError("Falta pytesseract en el entorno virtual") from exc

    command = shutil.which("tesseract")
    if command is None:
        from pathlib import Path
        for path in (
            Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
            Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
        ):
            if path.is_file():
                command = str(path)
                break
    if command is None:
        raise LocalOCRError("Tesseract no disponible en PATH o Program Files")
    pytesseract.pytesseract.tesseract_cmd = command
    try:
        languages = pytesseract.get_languages(config="")
    except Exception as exc:
        raise LocalOCRError("No se puede verificar el idioma de Tesseract") from exc
    if "spa" not in languages:
        raise LocalOCRError("Modelo spa de Tesseract no disponible")
    return pytesseract


def _ocr_pil_image(image: Any) -> str:
    if image.width * image.height > LOCAL_OCR_MAX_PIXELS:
        raise LocalOCRError("Imagen demasiado grande para OCR local")
    pytesseract = _local_tesseract()
    try:
        return clean_text(
            pytesseract.image_to_string(
                image.convert("RGB"),
                lang="spa+eng" if "eng" in pytesseract.get_languages(config="") else "spa",
                config="--psm 3",
                timeout=LOCAL_OCR_TIMEOUT_SECONDS,
            )
        )
    except Exception as exc:
        raise LocalOCRError(f"OCR local interrumpido: {type(exc).__name__}") from exc


def extract_document_text_local(content: bytes, file_name: str) -> tuple[str, str, str | None]:
    """Extrae texto nativo y OCR local selectivo (estado, texto, detalle).

    Se comparte entre CV, nominas y documentos firmados. No llama al proveedor LLM.
    Una incidencia de OCR debe impedir la sustitucion del indice previo, no marcar
    como completo un documento parcialmente leido.
    """
    if content.startswith(_IMAGE_MAGIC):
        from PIL import Image
        with Image.open(io.BytesIO(content)) as image:
            text = _ocr_pil_image(image)
        return ("ok" if len(text) >= 40 else "empty"), text, "ocr_local:imagen"

    if content[:4] != b"%PDF":
        status, text = extract_text(content, file_name)
        return status, text, None if status != "unsupported" else "formato no soportado"

    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    page_texts: list[str] = []
    pages_without_text: list[int] = []
    for index, page in enumerate(reader.pages):
        text = clean_text(page.extract_text() or "")
        page_texts.append(text)
        if len(text) < LOCAL_OCR_MIN_NATIVE_CHARS:
            pages_without_text.append(index)

    if len(pages_without_text) > LOCAL_OCR_MAX_SCANNED_PAGES:
        raise LocalOCRError(
            f"PDF con {len(pages_without_text)} paginas a reconocer; "
            f"limite OCR={LOCAL_OCR_MAX_SCANNED_PAGES}. No se indexa parcialmente"
        )

    if pages_without_text:
        import pymupdf
        from PIL import Image
        with pymupdf.open(stream=content, filetype="pdf") as pdf:
            if pdf.page_count != len(page_texts):
                raise LocalOCRError("Desajuste en numero de paginas PDF")
            for index in pages_without_text:
                page = pdf[index]
                # Calidad OCR y cota por pixeles para evitar cargas desproporcionadas.
                scale = min(1.9, (LOCAL_OCR_MAX_PIXELS / max(page.rect.width * page.rect.height, 1)) ** 0.5)
                pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), colorspace=pymupdf.csRGB, alpha=False)
                try:
                    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    page_texts[index] = _ocr_pil_image(image)
                finally:
                    del pix

    text = clean_text("\n".join(page_texts))
    detail = (f"ocr_local:pdf {len(pages_without_text)}/{len(page_texts)} paginas"
              if pages_without_text else None)
    return ("ok" if len(text) >= 40 else "empty"), text, detail


def extract_text(content: bytes, file_name: str) -> tuple[str, str]:
    """(estado, texto). Estados: ok · empty · unsupported.

    El tipo se decide por el contenido antes que por la extensión: hay imágenes
    guardadas como «.pdf».
    """
    if content.startswith(_IMAGE_MAGIC):
        return "unsupported", ""
    lower = file_name.lower()
    if content[:4] == b"%PDF":
        kind: str | None = "pdf"
    elif content[:2] == b"PK":
        kind = "docx"
    elif lower.endswith((".pdf", ".docx")):
        return "unsupported", ""  # la extensión no coincide con el contenido
    else:
        kind = None
    if kind == "pdf":
        from pypdf import PdfReader  # noqa: PLC0415

        reader = PdfReader(io.BytesIO(content))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif kind == "docx":
        import docx  # noqa: PLC0415

        document = docx.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs]
        for tbl in document.tables:
            parts.extend(cell.text for row in tbl.rows for cell in row.cells)
        text = "\n".join(parts)
    else:
        return "unsupported", ""
    text = clean_text(text)
    return ("ok", text) if len(text) >= 40 else ("empty", text)


def chunk_text(text: str) -> list[str]:
    chunks, start = [], 0
    while start < len(text):
        end = min(len(text), start + CHUNK_CHARS)
        if end < len(text):
            cut = text.rfind(" ", start + CHUNK_CHARS // 2, end)
            end = cut if cut > 0 else end
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return [c for c in chunks if c]


def normalize(vector: list[float]) -> list[float]:
    """Normaliza L2 antes de guardar/comparar — pgvector calcula la distancia coseno igual sin esto,
    pero se mantiene por consistencia con los vectores ya persistidos antes de la migración a pgvector."""
    array = np.asarray(vector, dtype=np.float32)
    norm = float(np.linalg.norm(array)) or 1.0
    return (array / norm).tolist()


class DocumentIndexer:
    """Indexador genérico para un dominio de documentos sensibles (nóminas, documentos firmados).

    Mismo flujo que ``CvIndexer`` (descarga → extracción/OCR → fragmentado → embeddings → persistencia),
    generalizado sobre ``HrDocumentText``/``HrDocumentChunk`` por ``domain`` + ``document_id`` (``sp_id``
    no es único entre bibliotecas distintas). Usa ``download_item_file`` (nunca ``download`` con una
    ruta, que puede contener un DNI).
    """

    def __init__(
        self,
        *,
        domain: str,
        list_title: str,
        model: type[Any],
        db: Database,
        client: SharePointClient,
        llm: LLMProvider,
        embedding_model: str,
        batch: int = 200,
        max_bytes: int = 15 * 1024 * 1024,
    ) -> None:
        self._domain = domain
        self._list_title = list_title
        self._model = model
        self._db = db
        self._client = client
        self._llm = llm
        self._model_name = embedding_model
        self._batch = batch
        self._max_bytes = max_bytes

    async def run(self, progress: Callable[[int, int], None] | None = None) -> IndexReport:
        report = IndexReport()
        model = self._model
        async with self._db.session() as db:
            alive = select(model.sp_id)
            removed = await db.execute(
                delete(HrDocumentText).where(
                    HrDocumentText.domain == self._domain, HrDocumentText.document_id.not_in(alive)
                )
            )
            await db.execute(
                delete(HrDocumentChunk).where(
                    HrDocumentChunk.domain == self._domain, HrDocumentChunk.document_id.not_in(alive)
                )
            )
            report.removed = int(getattr(removed, "rowcount", 0) or 0)
            stale = (
                select(model)
                .outerjoin(
                    HrDocumentText,
                    and_(HrDocumentText.domain == self._domain, HrDocumentText.document_id == model.sp_id),
                )
                .where(
                    or_(
                        HrDocumentText.document_id.is_(None),
                        HrDocumentText.source_modified != model.modified_at,
                        # Escaneados o imágenes que aún no han pasado por la lectura visual.
                        and_(
                            HrDocumentText.status.in_(("empty", "unsupported")),
                            or_(HrDocumentText.detail.is_(None), HrDocumentText.detail.not_like("ocr%")),
                        ),
                    )
                )
                .order_by(model.modified_at.desc())
            )
            pending = list((await db.scalars(stale)).all())
        report.pending = max(0, len(pending) - self._batch)
        for document in pending[: self._batch]:
            status = await self._index_one(document)
            report.processed += 1
            report.statuses[status] = report.statuses.get(status, 0) + 1
            report.indexed += status == "ok"
            report.failed += status == "error"
            if progress and (report.processed % 25 == 0 or report.processed == min(len(pending), self._batch)):
                progress(report.processed, min(len(pending), self._batch))
        log.info(
            "document_index.finished",
            domain=self._domain,
            processed=report.processed,
            indexed=report.indexed,
            failed=report.failed,
            removed=report.removed,
            pending=report.pending,
        )
        return report

    async def _index_one(self, document: Any) -> str:
        detail = None
        text = ""
        try:
            content = await self._client.download_item_file(self._list_title, document.sp_id, self._max_bytes)
            status, text, detail = await asyncio.to_thread(
                extract_document_text_local, content, document.file_name or ""
            )
        except LocalOCRError as exc:
            # Fallo seguro: no borrar vectores existentes ni intentar OCR multimodal.
            log.warning("document_index.local_ocr_failed", domain=self._domain,
                        document=document.sp_id, reason=str(exc))
            return "error"
        except SharePointError as exc:
            status, detail = ("too_large" if "demasiado grande" in str(exc) else "error"), str(exc)[:300]
        except Exception as exc:  # PDF o DOCX corrupto
            status, detail = "error", f"{type(exc).__name__}: {exc}"[:300]
        chunks = chunk_text(text) if status == "ok" else []
        vectors: list[list[float]] = []
        try:
            # Para NOMINAS, solo sale al proveedor una proyeccion de vocabulario
            # cerrado sin cifras ni datos personales. Texto y chunks ORIGINALES
            # siguen guardandose localmente para busquedas e informes autorizados.
            embedding_texts = (
                [project_payroll_embedding(chunk) for chunk in chunks]
                if self._domain == "nominas" else chunks
            )
            for start in range(0, len(embedding_texts), EMBED_BATCH):
                vectors.extend(await self._llm.embed(embedding_texts[start : start + EMBED_BATCH], self._model_name))
        except Exception as exc:
            log.warning("document_index.embed_failed", domain=self._domain, document=document.sp_id, error=str(exc))
            return "error"  # se reintentará en la siguiente ejecución (no se guarda el texto)
        # Marca version de privacidad en los NUEVOS documentos: los 8000 indices
        # anteriores permanecen intactos, sin fingir que fueron minimizados.
        if status == "ok" and self._domain == "nominas":
            detail = (
                f"{detail}; embeddings={PAYROLL_EMBEDDING_POLICY_VERSION}"
                if detail else f"embeddings={PAYROLL_EMBEDDING_POLICY_VERSION}"
            )
        try:
            await self._persist(document, status=status, text=text, detail=detail, chunks=chunks, vectors=vectors)
        except Exception as exc:  # un documento problemático nunca detiene la indexación del resto
            log.warning(
                "document_index.persist_failed", domain=self._domain, document=document.sp_id, error=str(exc)[:300]
            )
            await self._persist(document, status="error", text="", detail=type(exc).__name__, chunks=[], vectors=[])
            return "error"
        return status

    async def _persist(
        self,
        document: Any,
        *,
        status: str,
        text: str,
        detail: str | None,
        chunks: list[str],
        vectors: list[list[float]],
    ) -> None:
        async with self._db.session() as db:
            await db.execute(
                delete(HrDocumentChunk).where(
                    HrDocumentChunk.domain == self._domain, HrDocumentChunk.document_id == document.sp_id
                )
            )
            await db.execute(
                delete(HrDocumentText).where(
                    HrDocumentText.domain == self._domain, HrDocumentText.document_id == document.sp_id
                )
            )
            db.add(
                HrDocumentText(
                    domain=self._domain,
                    document_id=document.sp_id,
                    person_key=document.person_key,
                    source_modified=document.modified_at,
                    status=status,
                    text=text or None,
                    detail=detail,
                    extracted_at=datetime.now(UTC),
                )
            )
            db.add_all(
                HrDocumentChunk(
                    domain=self._domain,
                    document_id=document.sp_id,
                    person_key=document.person_key,
                    position=i,
                    text=chunk,
                    embedding=normalize(vector),
                )
                for i, (chunk, vector) in enumerate(zip(chunks, vectors, strict=True))
            )
