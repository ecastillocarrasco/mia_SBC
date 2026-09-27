"""Plantilla corporativa de informes (PDF, Excel y PowerPoint) a partir del modelo de documento."""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Any

from mia.domains.informes.document import Block, Document

NAVY = "1F4E78"
LIGHT = "F2F6FA"
FOOTER = "Sunset Beach Club · generado por mIA"


def _rows(block: Block) -> tuple[list[str], list[list[Any]]]:
    if block.kind == "table":
        return [str(c) for c in block.data.get("columns", [])], list(block.data.get("rows", []))
    return ["Concepto", "Valor"], [[k, v] for k, v in block.data.items()]


# ------------------------------------------------------------------ Excel
def to_xlsx(doc: Document, generated: datetime, logo: Path | None = None) -> bytes:
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.drawing.image import Image  # noqa: PLC0415
    from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    wb = Workbook()
    wb.remove(wb.active)
    for index, block in enumerate(doc.blocks or [Block("section", "Informe", {})], 1):
        ws = wb.create_sheet((block.title or f"Hoja {index}")[:28].replace("/", "-") or f"Hoja {index}")
        ws["A1"] = doc.title
        ws["A1"].font = Font(bold=True, size=14, color=NAVY)
        ws["A2"] = f"{block.title} · {FOOTER} el {generated:%d/%m/%Y %H:%M}"
        ws["A2"].font = Font(italic=True, size=9, color="666666")
        if logo and index == 1:
            image = Image(str(logo))
            image.height, image.width = 40, 40 * image.width / max(1, image.height)
            ws.add_image(image, "H1")
        columns, rows = _rows(block)
        header = 4
        for c, name in enumerate(columns, 1):
            cell = ws.cell(row=header, column=c, value=name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r, row in enumerate(rows, header + 1):
            for c, value in enumerate(row, 1):
                ws.cell(row=r, column=c, value=value)
        for c, name in enumerate(columns, 1):
            width = max([len(str(name))] + [len(str(row[c - 1] or "")) for row in rows[:500] if len(row) >= c]) + 2
            ws.column_dimensions[get_column_letter(c)].width = min(max(width, 8), 60)
        ws.freeze_panes = ws.cell(row=header + 1, column=1)
        if rows and columns:
            ws.auto_filter.ref = f"A{header}:{get_column_letter(len(columns))}{header + len(rows)}"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ PDF
def to_pdf(doc: Document, generated: datetime, logo: Path | None = None) -> bytes:
    from reportlab.lib import colors  # noqa: PLC0415
    from reportlab.lib.pagesizes import A4, landscape  # noqa: PLC0415
    from reportlab.lib.styles import getSampleStyleSheet  # noqa: PLC0415
    from reportlab.lib.units import cm  # noqa: PLC0415
    from reportlab.platypus import (  # noqa: PLC0415
        Image,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    widest = max((len(_rows(b)[0]) for b in doc.blocks), default=2)
    page = landscape(A4) if widest > 4 else A4
    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=page,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title=doc.title,
        author="mIA",
    )
    styles = getSampleStyleSheet()
    title_style = styles["Title"].clone("t", textColor=colors.HexColor(f"#{NAVY}"))
    h2 = styles["Heading2"].clone("h2", textColor=colors.HexColor(f"#{NAVY}"))
    cell = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
    head = cell.clone("head", textColor=colors.white, fontName="Helvetica-Bold")

    def para(value: Any, style: Any) -> Paragraph:
        text = "" if value is None else str(value)
        return Paragraph(text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), style)

    story: list[Any] = []
    if logo:
        story.append(Image(str(logo), width=3 * cm, height=1.5 * cm, kind="proportional"))
    story += [
        para(doc.title, title_style),
        para(f"{FOOTER} el {generated:%d/%m/%Y %H:%M}", styles["Italic"]),
        Spacer(1, 0.4 * cm),
    ]
    width = page[0] - 3 * cm
    for block in doc.blocks:
        columns, rows = _rows(block)
        if not rows:
            continue
        story.append(para(block.title, h2))
        grid = Table(
            [[para(c, head) for c in columns]] + [[para(v, cell) for v in row] for row in rows],
            colWidths=[width / max(1, len(columns))] * len(columns),
            repeatRows=1,
        )
        grid.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{LIGHT}")]),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C8D3DE")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story += [grid, Spacer(1, 0.4 * cm)]
    pdf.build(story)
    return buffer.getvalue()


# ------------------------------------------------------------------ PowerPoint
def to_pptx(doc: Document, generated: datetime, logo: Path | None = None) -> bytes:
    from pptx import Presentation  # noqa: PLC0415
    from pptx.dml.color import RGBColor  # noqa: PLC0415
    from pptx.util import Inches, Pt  # noqa: PLC0415

    navy = RGBColor.from_string(NAVY)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def header(slide: Any, text: str) -> None:
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.9))
        bar.fill.solid()
        bar.fill.fore_color.rgb = navy
        bar.line.fill.background()
        box = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(11.5), Inches(0.6)).text_frame
        box.text = text
        box.paragraphs[0].runs[0].font.size = Pt(24)
        box.paragraphs[0].runs[0].font.bold = True
        box.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        foot = slide.shapes.add_textbox(Inches(0.4), Inches(7.05), Inches(12), Inches(0.3)).text_frame
        foot.text = f"{FOOTER} · {generated:%d/%m/%Y}"
        foot.paragraphs[0].runs[0].font.size = Pt(9)
        if logo:
            slide.shapes.add_picture(str(logo), Inches(12.2), Inches(0.1), height=Inches(0.7))

    cover = prs.slides.add_slide(blank)
    header(cover, "Sunset Beach Club")
    title = cover.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.5)).text_frame
    title.text = doc.title
    title.paragraphs[0].runs[0].font.size = Pt(36)
    title.paragraphs[0].runs[0].font.color.rgb = navy
    for block in doc.blocks:
        if block.kind == "kpis":
            _kpi_slide(prs, blank, header, block)
        else:
            _table_slides(prs, blank, header, block)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _kpi_slide(prs: Any, layout: Any, header: Any, block: Block) -> None:
    from pptx.dml.color import RGBColor  # noqa: PLC0415
    from pptx.util import Inches, Pt  # noqa: PLC0415

    navy = RGBColor.from_string(NAVY)
    slide = prs.slides.add_slide(layout)
    header(slide, block.title)
    for i, (label, value) in enumerate(list(block.data.items())[:12]):
        x, y = Inches(0.5 + (i % 4) * 3.15), Inches(1.3 + (i // 4) * 1.8)
        card = slide.shapes.add_shape(1, x, y, Inches(2.95), Inches(1.6))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor.from_string(LIGHT)
        card.line.color.rgb = navy
        tf = card.text_frame
        tf.text = str(value)[:40]
        tf.paragraphs[0].runs[0].font.size = Pt(22 if len(str(value)) < 12 else 12)
        tf.paragraphs[0].runs[0].font.color.rgb = navy
        tf.add_paragraph().text = str(label)
        tf.paragraphs[1].runs[0].font.size = Pt(11)


def _table_slides(prs: Any, layout: Any, header: Any, block: Block) -> None:
    from pptx.util import Inches, Pt  # noqa: PLC0415

    columns, rows = _rows(block)
    for start in range(0, len(rows), 12):
        chunk = rows[start : start + 12]
        slide = prs.slides.add_slide(layout)
        header(slide, block.title if start == 0 else f"{block.title} (cont.)")
        shape = slide.shapes.add_table(
            len(chunk) + 1,
            len(columns),
            Inches(0.4),
            Inches(1.2),
            Inches(12.5),
            Inches(0.4) * (len(chunk) + 1),
        )
        grid = shape.table
        for c, name in enumerate(columns):
            grid.cell(0, c).text = str(name)
        for r, row in enumerate(chunk, 1):
            for c in range(len(columns)):
                grid.cell(r, c).text = "" if c >= len(row) or row[c] is None else str(row[c])[:120]
        for r in range(len(chunk) + 1):
            for c in range(len(columns)):
                for paragraph in grid.cell(r, c).text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(11 if r == 0 else 10)


RENDERERS = {"pdf": to_pdf, "xlsx": to_xlsx, "pptx": to_pptx}
