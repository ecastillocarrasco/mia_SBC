"""Indexación de CVs: descarga, extracción de texto, fragmentos y embeddings.

Se ejecuta tras la sincronización. Solo procesa CVs nuevos o modificados (por
lotes), así que el coste de embeddings se paga una vez por CV.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import UTC, datetime

import numpy as np
from sqlalchemy import delete, or_, select

from mia.data import CvChunk, CvDocument, CvText, Database
from mia.llm import LLMProvider
from mia.observability import get_logger
from mia.sources.sharepoint import SharePointClient, SharePointError

log = get_logger(__name__)

CHUNK_CHARS = 1200
CHUNK_OVERLAP = 150
EMBED_BATCH = 64


@dataclass(slots=True)
class IndexReport:
    processed: int = 0
    indexed: int = 0
    failed: int = 0
    removed: int = 0
    pending: int = 0
    statuses: dict[str, int] = field(default_factory=dict)


def extract_text(content: bytes, file_name: str) -> tuple[str, str]:
    """(estado, texto). Estados: ok · empty · unsupported."""
    lower = file_name.lower()
    kind = "pdf" if lower.endswith(".pdf") or content[:4] == b"%PDF" else None
    if kind is None and (lower.endswith(".docx") or content[:2] == b"PK"):
        kind = "docx"
    if kind == "pdf":
        from pypdf import PdfReader  # noqa: PLC0415

        reader = PdfReader(io.BytesIO(content))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif kind == "docx":
        import docx  # noqa: PLC0415

        document = docx.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs]
        for tbl in document.tables:
            parts.extend(cell.text for row in tbl.rows for cell in row.cells)
        text = "\n".join(parts)
    else:
        return "unsupported", ""
    text = " ".join(text.split())
    return ("ok", text) if len(text) >= 40 else ("empty", text)


def chunk_text(text: str) -> list[str]:
    chunks, start = [], 0
    while start < len(text):
        end = min(len(text), start + CHUNK_CHARS)
        if end < len(text):
            cut = text.rfind(" ", start + CHUNK_CHARS // 2, end)
            end = cut if cut > 0 else end
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return [c for c in chunks if c]


def pack(vector: list[float]) -> bytes:
    array = np.asarray(vector, dtype=np.float32)
    norm = float(np.linalg.norm(array)) or 1.0
    return (array / norm).astype(np.float32).tobytes()


class CvIndexer:
    def __init__(
        self,
        *,
        db: Database,
        client: SharePointClient,
        llm: LLMProvider,
        embedding_model: str,
        batch: int = 200,
        max_bytes: int = 15 * 1024 * 1024,
    ) -> None:
        self._db = db
        self._client = client
        self._llm = llm
        self._model = embedding_model
        self._batch = batch
        self._max_bytes = max_bytes

    async def run(self) -> IndexReport:
        report = IndexReport()
        async with self._db.session() as db:
            alive = select(CvDocument.sp_id)
            removed = await db.execute(delete(CvText).where(CvText.document_id.not_in(alive)))
            await db.execute(delete(CvChunk).where(CvChunk.document_id.not_in(alive)))
            report.removed = int(getattr(removed, "rowcount", 0) or 0)
            stale = (
                select(CvDocument)
                .outerjoin(CvText, CvText.document_id == CvDocument.sp_id)
                .where(or_(CvText.document_id.is_(None), CvText.source_modified != CvDocument.modified_at))
                .order_by(CvDocument.modified_at.desc())
            )
            pending = list((await db.scalars(stale)).all())
        report.pending = max(0, len(pending) - self._batch)
        for document in pending[: self._batch]:
            status = await self._index_one(document)
            report.processed += 1
            report.statuses[status] = report.statuses.get(status, 0) + 1
            report.indexed += status == "ok"
            report.failed += status == "error"
        log.info(
            "cv_index.finished",
            processed=report.processed,
            indexed=report.indexed,
            failed=report.failed,
            removed=report.removed,
            pending=report.pending,
        )
        return report

    async def _index_one(self, document: CvDocument) -> str:
        detail = None
        text = ""
        try:
            content = await self._client.download(document.file_ref, self._max_bytes)
            status, text = extract_text(content, document.file_name)
        except SharePointError as exc:
            status, detail = ("too_large" if "demasiado grande" in str(exc) else "error"), str(exc)[:300]
        except Exception as exc:  # PDF o DOCX corrupto
            status, detail = "error", f"{type(exc).__name__}: {exc}"[:300]
        chunks = chunk_text(text) if status == "ok" else []
        vectors: list[list[float]] = []
        try:
            for start in range(0, len(chunks), EMBED_BATCH):
                vectors.extend(await self._llm.embed(chunks[start : start + EMBED_BATCH], self._model))
        except Exception as exc:
            log.warning("cv_index.embed_failed", document=document.sp_id, error=str(exc))
            return "error"  # se reintentará en la siguiente ejecución (no se guarda el texto)
        async with self._db.session() as db:
            await db.execute(delete(CvChunk).where(CvChunk.document_id == document.sp_id))
            await db.execute(delete(CvText).where(CvText.document_id == document.sp_id))
            db.add(
                CvText(
                    document_id=document.sp_id,
                    candidate_id=document.candidate_id,
                    source_modified=document.modified_at,
                    status=status,
                    text=text or None,
                    detail=detail,
                    extracted_at=datetime.now(UTC),
                )
            )
            db.add_all(
                CvChunk(
                    document_id=document.sp_id,
                    candidate_id=document.candidate_id,
                    position=i,
                    text=chunk,
                    embedding=pack(vector),
                )
                for i, (chunk, vector) in enumerate(zip(chunks, vectors, strict=True))
            )
        return status
