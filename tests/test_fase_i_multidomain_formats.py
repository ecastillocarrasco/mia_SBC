"""Fase I, punto 2: "completar cobertura multidominio" — extiende la correspondencia exacta
pantalla↔documento (ya probada para nóminas en `test_fase_i_integral_validation.py`) a Turnos,
Vacaciones, Plantilla y Ausencias, en los 4 formatos soportados (PDF/XLSX/PPTX/HTML).

`generar_informe` (`src/mia/modules/rrhh/informes/tools.py`) es genérico: construye el documento a
partir de `LAST_TABLE` (lo que sea que la última herramienta haya recordado con `remember_table`), no
hay una herramienta de documento distinta por dominio. Por eso la matriz real a comprobar es "dominio ×
formato", no "dominio × capacidad documental propia" — no existe tal cosa como un dominio sin capacidad
documental en este sistema; si alguna combinación no funcionara, sería un fallo de `generar_informe`
o de un `RENDERERS[formato]`, no la ausencia de una función dedicada.

Usa los datos sembrados reales de `tests/hr_seed.py::seed()` (Ana Ruiz Sanz: vacaciones 3-7 agosto,
descanso 8-9, trabajo 10-14; Eva Sol: trabajo 3-4, baja 5-6; TODAY = 2026-08-05), a través del punto de
entrada real del 2D (`/api/mia/chat`).
"""

from __future__ import annotations

import io
import json
from typing import Any

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from mia.llm.fake import final_answer, tool_call
from tests.hr_seed import seed

FORMATS = ["pdf", "xlsx", "pptx", "html"]

# (etiqueta de dominio, herramienta, argumentos, valor que debe aparecer igual en pantalla y documento)
DOMAIN_CASES = [
    ("ausencias", "ausencias", {"nombre": "Eva Sol", "desde": "2026-08-03", "hasta": "2026-08-06"}, "Eva Sol"),
    ("turnos", "turnos_persona", {"nombre": "Ana Ruiz Sanz", "desde": "2026-08-03", "hasta": "2026-08-10"}, "Vacaciones"),
    ("vacaciones", "vacaciones", {"desde": "2026-08-01", "hasta": "2026-08-31"}, "Ana Ruiz Sanz"),
    ("plantilla", "plantilla", {}, "Ana Ruiz Sanz"),
]


def _download_action(body: dict) -> dict:
    for action in body.get("actions") or []:
        if action.get("type") == "download":
            return action
    raise AssertionError(f"no se encontró una acción de descarga en actions: {body.get('actions')}")


def _extract_text(formato: str, content: bytes) -> str:
    if formato == "pdf":
        return "".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    if formato == "xlsx":
        wb = load_workbook(io.BytesIO(content))
        return " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)
    if formato == "pptx":
        parts: list[str] = []
        for slide in Presentation(io.BytesIO(content)).slides:
            for shape in slide.shapes:
                if shape.has_table:
                    parts += [cell.text for row in shape.table.rows for cell in row.cells]
                elif shape.has_text_frame:
                    parts.append(shape.text_frame.text)
        return " ".join(parts)
    return content.decode("utf-8")


@pytest.mark.parametrize("formato", FORMATS)
@pytest.mark.parametrize("domain_label,tool_name,tool_args,expected_value", DOMAIN_CASES, ids=[c[0] for c in DOMAIN_CASES])
async def test_domain_table_matches_the_downloaded_document(
    client: Any, services: Any, script: list,
    domain_label: str, tool_name: str, tool_args: dict, expected_value: str, formato: str,
) -> None:
    await seed(services.db)

    def scripted(req):
        tool_rounds = [m for m in req.messages if m.role == "tool"]
        if not tool_rounds:
            return tool_call(tool_name, tool_args)
        if len(tool_rounds) == 1:
            return tool_call("generar_informe", {"formato": formato})
        return final_answer(f"Aquí tienes el informe de {domain_label}.", "Listo.")

    script[0] = scripted
    response = await client.post("/api/mia/chat", json={"message": f"Informe de {domain_label}"})
    assert response.status_code == 200
    body = response.json()
    data = body["data"]

    table = data.get("table")
    assert table and table.get("rows"), f"[{domain_label}] no llegó como tabla al mundo 2D: {data}"
    screen_text = json.dumps(table["rows"], ensure_ascii=False)
    assert expected_value in screen_text, f"[{domain_label}] la pantalla no muestra {expected_value!r}: {screen_text}"

    action = _download_action(body)
    download = await client.get(action["url"])
    assert download.status_code == 200, f"[{domain_label}/{formato}] descarga fallida: {download.status_code}"

    doc_text = _extract_text(formato, download.content)
    assert expected_value in doc_text, (
        f"[{domain_label}/{formato}] {expected_value!r} aparece en pantalla pero no en el documento "
        f"descargado — posible pérdida o sustitución de datos en generar_informe/RENDERERS"
    )
