"""Regresión de la incidencia real (trace ``mia-20261012-115403-bad97707``): "exporta a pdf el
informe anterior" terminaba con respuesta HTTP/operativa "ok" pero sin presentaciones ni archivo
(``data``/``actions``/``presentations`` vacíos) — nunca una excepción visible para el usuario, solo
un resultado vacío en silencio.

Causa raíz confirmada (reproducida de verdad, no suposición — ver `tests/test_incident_tool_call_
error_repro.py` para el detalle de la investigación): cuando el modelo pide "un informe completo"
y el router elige `generar_informe_ejecutivo` (en vez de `analisis_economico_nominas` directamente),
la redirección real de la sección 15 SÍ genera el PDF económico correcto en ese turno — pero la tabla
real del análisis (`eco_result.presentation`, con KPIs/gráficos/departamentos) solo vivía en una
variable local de `_redirect_to_economic_analysis()`. `remember_table()` descarta a propósito el
resultado `{"type":"download"}` de ese mismo informe (ya generado, nada que re-exportar) — así que la
tabla real NUNCA llegaba a `session_state`. Una petición posterior en OTRO turno ("exporta de nuevo",
"ahora en Excel"…) no encontraba ningún resultado que exportar.

Corrección: `_redirect_to_economic_analysis()` (`src/mia/modules/rrhh/informes/tools.py`) ahora llama
a `remember_table()` con la tabla REAL del análisis económico antes de renderizar el informe ejecutivo
— el mismo patrón que ya usa `run_informe_economico_nominas` (Job Manager, Fase G) por la misma razón.

Esta prueba reproduce la conversación de aceptación completa pedida (5 turnos reales, misma sesión,
sin truco): análisis → PDF → PowerPoint → Excel → HTML, todos sobre el MISMO resultado autorizado, con
datos sintéticos controlados (no el importe real de la incidencia)."""

from __future__ import annotations

import io

import httpx
import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from mia.connectors.sharepoint import SharePointClient, StaticTokenProvider
from mia.data import Payslip
from mia.llm.base import LLMCompleted, LLMRequest, TextDelta, ToolCall, ToolCallRequested
from mia.llm.fake import FakeLLMProvider, final_answer, tool_call
from mia.main import create_app
from tests.conftest import make_settings
from tests.dbutil import fresh_database
from tests.hr_seed import _emp
from tests.test_phase4b import _pdf
from tests.test_report_routing_semantic_validation import _payslip_text_with_ss

_SYNTHETIC_DEPARTMENTS: dict[str, dict[str, float]] = {
    "Cocina": {f"pk-c{i}": 1600.0 + i * 15 for i in range(1, 10)},
    "Recepción": {f"pk-r{i}": 1900.0 + i * 12 for i in range(1, 8)},
    "Pisos": {f"pk-p{i}": 1500.0 + i * 10 for i in range(1, 12)},
    "Mantenimiento": {"pk-m1": 1700.0, "pk-m2": 1750.0},  # 2 personas: oculto por NOM-005
}


def _extract_text(formato: str, content: bytes) -> str:
    if formato == "pdf":
        return "".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    if formato == "xlsx":
        # A diferencia de PDF/PPTX/HTML, una celda numérica en Excel guarda el VALOR nativo (p. ej.
        # 2392.19) más un number_format (p. ej. '#,##0.00 "€"') — el texto "2.392,19 €" es solo cómo
        # Excel lo MUESTRA, nunca un literal en el archivo. Para poder comparar contra el texto
        # extraído de los otros formatos, se reconstruye aquí esa misma representación con formato
        # español para las celdas que llevan ese number_format, en vez de buscar el texto formateado
        # como si fuera a aparecer tal cual (nunca podría encontrarlo, aunque el dato fuera correcto).
        wb = load_workbook(io.BytesIO(content))
        parts: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    parts.append(str(c.value))
                    if isinstance(c.value, int | float) and not isinstance(c.value, bool) and "€" in (c.number_format or ""):
                        parts.append(f"{c.value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") + " €")
        return " ".join(parts)
    if formato == "pptx":
        parts: list[str] = []
        for slide in Presentation(io.BytesIO(content)).slides:
            for shape in slide.shapes:
                if shape.has_table:
                    parts += [cell.text for row in shape.table.rows for cell in row.cells]
                elif shape.has_text_frame:
                    parts.append(shape.text_frame.text)
        return " ".join(parts)
    return content.decode("utf-8")


def _download_action(body: dict) -> dict:
    for action in body.get("actions") or []:
        if action.get("type") == "download":
            return action
    raise AssertionError(f"no se encontró una acción de descarga: {body.get('actions')}")


@pytest.fixture
async def acceptance_app(tmp_path):
    db = await fresh_database(tmp_path, "incident_acceptance.db")
    async with db.session() as s:
        sp_id = 6000
        for dept, people in _SYNTHETIC_DEPARTMENTS.items():
            position = "Cocinero" if dept == "Cocina" else "Empleado"
            for pk in people:
                sp_id += 1
                s.add(_emp(sp_id, f"Empleado {pk}", dept, position=position, person_key=pk))
                s.add(Payslip(sp_id=sp_id, person_key=pk, year=2026, month=8, period="2026-08", special=False))

    net_by_sp_id: dict[int, float] = {}
    sp_id = 6000
    for people in _SYNTHETIC_DEPARTMENTS.values():
        for net in people.values():
            sp_id += 1
            net_by_sp_id[sp_id] = net

    def files(request: httpx.Request) -> httpx.Response:
        import re  # noqa: PLC0415

        match = re.search(r"items\((\d+)\)/File", request.url.path)
        if not match:
            return httpx.Response(404)
        return httpx.Response(200, content=_pdf(_payslip_text_with_ss(net_by_sp_id[int(match.group(1))])))

    sp = SharePointClient(
        site_url="https://x/sites/rrhh", tokens=StaticTokenProvider(), transport=httpx.MockTransport(files)
    )

    # Guion de los 5 turnos de aceptación: turno 1 pasa por el router "informe completo" ->
    # generar_informe_ejecutivo (el camino que disparó la incidencia real, no analisis_economico_
    # nominas directamente — así la prueba ejercita de verdad la corrección, no la rodea).
    turn = {"n": 0}
    FORMATS_BY_TURN = {2: "pdf", 3: "pptx", 4: "xlsx", 5: "html"}

    def scripted(req: LLMRequest):
        # compose_for_request() (motor documental, content_composer.py) hace su PROPIA llamada al LLM
        # para elegir qué bloques incluir — independiente del bucle principal de herramientas de este
        # guion, y la dispara el turno 1 (no encaja en ninguna de las frases de "solo reexporta" que
        # omiten esa llamada). Reconocerla por su response_schema propio (selected_indices/unavailable)
        # evita que este guion la confunda con el inicio de un turno nuevo y corrompa turn["n"].
        schema_props = (req.response_schema or {}).get("properties") or {}
        if "selected_indices" in schema_props:
            import json as _json

            inventory = _json.loads(req.messages[0].content)["bloques"]
            decision = _json.dumps({"selected_indices": [b["index"] for b in inventory], "unavailable": []})
            return [TextDelta(decision), LLMCompleted()]

        tool_rounds = [m for m in req.messages if m.role == "tool"]
        if not tool_rounds:
            turn["n"] += 1
            if turn["n"] == 1:
                return [
                    ToolCallRequested(
                        ToolCall("call_1", "generar_informe_ejecutivo", '{"periodo": "agosto de 2026", "formato": "pdf"}')
                    ),
                ]
            formato = FORMATS_BY_TURN[turn["n"]]
            return tool_call("generar_informe", {"formato": formato, "incluir_detalle": True, "incluir_grafico": True})
        return final_answer("Aquí tienes el resultado.", "Listo.")

    llm = FakeLLMProvider(scripted)
    settings = make_settings(tmp_path, database_url=f"sqlite+aiosqlite:///{tmp_path / 'incident_acceptance.db'}")

    from mia.api.services import Services
    from mia.conversation import ConversationEngine, Persona
    from mia.modules.registry import build_registry

    services = Services.build(settings, llm=llm, db=db)
    registry = build_registry(
        db, settings.data.stale_after_minutes, settings=settings, llm=llm, sharepoint=sp, enabled=settings.domains.enabled
    )
    services.registry = registry
    services.sharepoint = sp
    services.engine = ConversationEngine(
        llm=llm, tools=registry, db=db, persona=Persona(), timezone=settings.timezone,
        history_turns=settings.history_turns, max_tool_rounds=settings.llm.max_tool_rounds,
        turn_budget_seconds=settings.llm.turn_budget_seconds, max_output_tokens=settings.llm.max_output_tokens,
    )
    app = create_app(settings, services=services)
    yield app
    await sp.aclose()
    await db.dispose()


async def test_acceptance_conversation_all_five_turns_export_the_same_authorized_analysis(acceptance_app) -> None:
    """Los 5 turnos de aceptación pedidos. Los 4 documentos (PDF/PPTX/XLSX/HTML) deben proceder del
    MISMO análisis autorizado: mismos importes, mismo periodo, mismos departamentos, mismos
    porcentajes, misma advertencia NOM-005 — nunca un recálculo paralelo ni un documento distinto."""
    messages = [
        "Genera un informe completo de comparativa de costes en nóminas por departamento en agosto de 2026",
        "Exporta a PDF el informe anterior",
        "PowerPoint",
        "Ahora en Excel",
        "También en HTML",
    ]
    texts_by_format: dict[str, str] = {}
    session_id: str | None = None

    async with acceptance_app.router.lifespan_context(acceptance_app):
        transport = httpx.ASGITransport(app=acceptance_app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            for i, message in enumerate(messages, start=1):
                payload = {"message": message}
                if session_id:
                    payload["context"] = {"session_id": session_id}
                response = await client.post("/api/mia/chat", json=payload)
                assert response.status_code == 200, f"turno {i} ({message!r}): HTTP {response.status_code}"
                body = response.json()
                assert body["status"] == "ok", f"turno {i} ({message!r}): status={body['status']}"
                session_id = body["context"]["session_id"]

                if i == 1:
                    continue  # el primer turno genera el PDF "ejecutivo" (ya cubierto); los formatos
                    # pedidos explícitamente empiezan en el turno 2.

                action = _download_action(body)
                assert body["presentations"], f"turno {i} ({message!r}): sin presentaciones — el síntoma exacto de la incidencia"
                download = await client.get(action["url"])
                assert download.status_code == 200
                formato = {2: "pdf", 3: "pptx", 4: "xlsx", 5: "html"}[i]
                texts_by_format[formato] = _extract_text(formato, download.content)

    assert set(texts_by_format) == {"pdf", "pptx", "xlsx", "html"}

    # Correspondencia exacta entre los 4 documentos: mismos departamentos, mismo periodo, misma
    # advertencia de privacidad — ninguno es un recálculo distinto ni un documento ajeno.
    for formato, text in texts_by_format.items():
        assert "Cocina" in text, f"{formato}: falta Cocina"
        assert "Recepción" in text, f"{formato}: falta Recepción"
        assert "Pisos" in text, f"{formato}: falta Pisos"
        assert "NOM-005" in text, f"{formato}: falta la advertencia de grupo oculto (Mantenimiento, 2 personas)"
        assert "1.700,00" not in text and "1.750,00" not in text, (
            f"{formato}: se filtró un importe individual del grupo oculto (Mantenimiento)"
        )

    # Las cifras por departamento deben coincidir EXACTAMENTE entre todos los formatos — se extraen
    # del propio PDF (ya verificado arriba que contiene los datos reales) y se buscan literalmente en
    # los otros tres, en vez de recalcularlas aquí de forma independiente.
    import re as _re

    pdf_amounts = sorted(set(_re.findall(r"\d{1,3}(?:\.\d{3})*,\d{2}\s?€", texts_by_format["pdf"])))
    assert pdf_amounts, "no se encontró ninguna cifra en euros en el PDF para comparar"
    for formato in ("pptx", "xlsx", "html"):
        missing = [amt for amt in pdf_amounts if amt.replace(" ", "").rstrip("€").rstrip() not in texts_by_format[formato].replace(" ", "")]
        assert not missing, f"{formato}: cifras presentes en el PDF pero no en este formato: {missing}"
