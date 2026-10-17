"""Contrato corporativo: todo informe lleva el logotipo SBC en la cabecera."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import fitz
import openpyxl
import pytest
from pptx import Presentation

from mia.assets import resolve_logo
from mia.modules.rrhh.informes.document import Block, Document
from mia.modules.rrhh.informes.render import to_pdf, to_pptx, to_xlsx
from mia.modules.rrhh.informes.render_html import to_html

NOW = datetime(2026, 10, 14)


def fixture_document():
    # Varias páginas, diapositivas y pestañas, además de la portada PPTX.
    rows = [[f"Departamento {i}", float(i)] for i in range(95)]
    return Document("Informe corporativo de prueba", [
        Block("table", "Detalle 1", {"columns": ["Área", "Coste"], "rows": rows}),
        Block("table", "Detalle 2", {"columns": ["Área", "Coste"], "rows": rows}),
    ])


def test_logo_present_across_all_pages_sheets_slides_and_html():
    doc = fixture_document()
    pdf = fitz.open(stream=to_pdf(doc, NOW), filetype="pdf")
    assert len(pdf) > 1
    for page in pdf:
        # The logo is a real embedded image on every page (not text-only branding).
        images = page.get_images(full=True)
        assert images, f"Falta imagen en la página PDF {page.number + 1}"
        # it must lie within the upper page header region.
        assert any(rect.y0 < page.rect.height * 0.18
                   for img in images for rect in page.get_image_rects(img[0]))

    wb = openpyxl.load_workbook(BytesIO(to_xlsx(doc, NOW)))
    assert len(wb.worksheets) == 2
    assert all(len(ws._images) == 1 and ws._images[0].anchor._from.row == 0
               for ws in wb.worksheets)

    ppt = Presentation(BytesIO(to_pptx(doc, NOW)))
    assert len(ppt.slides) > 2
    for slide in ppt.slides:
        assert any(shape.shape_type == 13 and shape.top < 1100000
                   for shape in slide.shapes), "Diapositiva sin logo en cabecera"

    html = to_html(doc, NOW).decode("utf-8")
    assert '<header class="header">' in html
    assert 'alt="Logotipo Sunset Beach Club"' in html
    assert 'src="data:image/png;base64,' in html


def test_missing_packaged_logo_blocks_export(monkeypatch, tmp_path):
    # Ni siquiera una exportación llamada directamente debe funcionar sin marca.
    import mia.assets
    monkeypatch.setattr(mia.assets, "LOGO", tmp_path / "no_existe.png")
    document = Document("Informe", [Block("section", "Resumen", {"dato": 1})])
    for exporter in (to_pdf, to_pptx, to_xlsx, to_html):
        with pytest.raises(FileNotFoundError, match="logotipo corporativo obligatorio"):
            exporter(document, NOW)


def test_invalid_override_uses_packaged_logo():
    assert resolve_logo(Path("/ruta/inexistente/otro-logo.png")).is_file()
