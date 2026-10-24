"""Ajuste 4 de la aprobación condicionada del Bloque I — dominio Informes: contrato de seguridad y
prueba de integración real.

El Advanced Intelligence & Reporting Engine (Bloque II, con marca Sunset Beach Club obligatoria en
PDF/XLSX/PPTX) todavía no existe: sus propias pruebas de extremo a extremo se ejecutarán durante ese
bloque. Pero el generador de informes de RRHH YA EXISTE hoy (``mia.modules.rrhh.informes``) y es
exactamente sobre este generador donde se puede, y se debe, fijar ahora el contrato de seguridad que el
motor de Bloque II deberá heredar sin reinventarlo ni debilitarlo:

    CONTRATO: el documento final (PDF/XLSX/PPTX) se renderiza en el backend directamente a partir de la
    última presentación de herramienta guardada en el estado de la conversación (``LAST_TABLE``, ver
    ``remember_table()`` en ``mia/conversation/tools.py``) — datos que llegan a ``ToolContext.session_state``
    ANTES de que el Privacy Gateway sanitice nada hacia el LLM (ver ``ConversationEngine.run()``, línea
    donde se llama a ``remember_table`` justo después de ejecutar la herramienta, no después de hablar
    con el proveedor). El LLM nunca ve, ni necesita ver, el dato real para que el informe sea correcto.

Esta prueba lo demuestra de extremo a extremo, no solo lo afirma: genera un ranking salarial real
(nombres ligados a importes — el caso más sensible, NOM-004/005), confirma que el LLM solo vio
pseudo-tokens, y abre el .xlsx generado de verdad para comprobar que SÍ contiene el nombre y el importe
reales. Es la garantía exacta que exige el ajuste 1 de la aprobación condicionada: "el backend pueda
generar los documentos finales utilizando los datos originales autorizados, sin depender de que el LLM
conozca las cifras individuales".

Dominio Correo: ver ``tests/test_privacy_gateway_mail_integration.py`` para el equivalente de este
contrato aplicado a correos del buzón personal autenticado.
"""

from __future__ import annotations

import io
import json
import re
from pathlib import Path
from types import SimpleNamespace

import httpx
from openpyxl import load_workbook
from sqlalchemy import select

from mia.auth import UserContext
from mia.connectors.sharepoint import SharePointClient, StaticTokenProvider
from mia.conversation import ConversationEngine
from mia.conversation.events import ResultEvent
from mia.data import Employee, Payslip, Report
from mia.llm.base import LLMRequest
from mia.llm.fake import FakeLLMProvider, final_answer, tool_call
from mia.modules.registry import build_registry
from mia.privacy_gateway.gateway import PrivacyGateway
from mia.settings import PrivacyGatewaySettings
from tests.dbutil import fresh_database
from tests.test_phase4b import _pdf

ADMIN = UserContext("u", "Enrique", roles=frozenset({"mia.admin"}))


def _stub_settings(reports_dir: Path) -> SimpleNamespace:
    return SimpleNamespace(
        reports=SimpleNamespace(directory=str(reports_dir), ttl_hours=24, logo_path=None),
        data=SimpleNamespace(content_enabled=True),
        mail=None,
    )


async def test_generated_xlsx_report_contains_real_data_while_llm_only_saw_pseudo_tokens(tmp_path: Path) -> None:
    names_and_nets = [("Marta Ruiz", "pk-401", 1700.0), ("Carlos López", "pk-402", 1600.0), ("Sofía Pérez", "pk-403", 1500.0)]
    db = await fresh_database(tmp_path, "privacy_informes_contract.db")
    async with db.session() as s:
        for n, (name, pk, _net) in enumerate(names_and_nets, start=401):
            s.add(
                Employee(
                    sp_id=n, name=name, name_key=name.lower(), active=True, department="Cocina",
                    department_key="cocina", position="Cocinero", person_key=pk,
                )
            )
            s.add(Payslip(sp_id=n, person_key=pk, year=2026, month=8, period="2026-08", special=False))

    net_by_spid = {401 + i: net for i, (_, _, net) in enumerate(names_and_nets)}

    def files(request: httpx.Request) -> httpx.Response:
        match = re.search(r"items\((\d+)\)/File", request.url.path)
        net = net_by_spid[int(match.group(1))] if match else 0.0
        gross = net + 330.15
        fmt = lambda v: f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")  # noqa: E731
        text = (
            "EMPRESA SUNSET BEACH CLUB\nMensual - Agosto 2026\n"
            "REM. TOTAL P.P. EXTRAS BASE C.C. BASE A.T. Y DES BASE I.R.P.F. T. DEVENGADO T. A DEDUCIR\n"
            f"{fmt(gross - 266.67)} 266,67 {fmt(gross)} {fmt(gross)} {fmt(gross)} {fmt(gross)} 330,15\n"
            f"Tributación I.R.P.F. (11,50 %) 214,67\n{fmt(net)} Líquido a percibir\n"
        )
        return httpx.Response(200, content=_pdf(text))

    sp = SharePointClient(site_url="https://x/sites/rrhh", tokens=StaticTokenProvider(), transport=httpx.MockTransport(files))
    settings = _stub_settings(tmp_path / "informes_generados")
    registry = build_registry(db, settings=settings, sharepoint=sp)

    def script(request: LLMRequest) -> list[object]:
        last = request.messages[-1]
        if last.role == "tool" and last.tool_call_id == "c_ranking":
            return tool_call(
                "generar_informe",
                {"formato": "xlsx", "incluir_detalle": True, "incluir_grafico": False},
                call_id="c_informe",
            )
        if last.role == "tool" and last.tool_call_id == "c_informe":
            return final_answer(texto=last.content, habla="Informe generado.", estado_avatar="sonriendo")
        return tool_call("ranking_salarial", {"departamento": "Cocina", "cuantos": 3}, call_id="c_ranking")

    fake = FakeLLMProvider(script=script)
    gw_settings = PrivacyGatewaySettings(enabled=True, vault_key="clave-de-prueba-no-real", vault_ttl_seconds=60)
    gateway = PrivacyGateway(fake, gw_settings, provider_name="test", model="test-model")
    engine = ConversationEngine(llm=gateway, tools=registry, db=db, max_tool_rounds=4)

    result = None
    async for event in engine.run(
        message="Genera un informe Excel del ranking salarial de Cocina", user=ADMIN, session_id=None
    ):
        if isinstance(event, ResultEvent):
            result = event.result
    await sp.aclose()
    assert result is not None
    assert result.status == "ok"

    # 1) Lo que de verdad salió hacia el "proveedor" (el LLM) nunca llevó el dato real.
    outgoing = json.dumps([m.content for req in fake.requests for m in req.messages], ensure_ascii=False)
    assert "Marta Ruiz" not in outgoing
    assert "Carlos López" not in outgoing
    assert "[PERSON_" in outgoing

    # 2) El archivo .xlsx generado de verdad por el backend SÍ contiene el dato real: el informe no
    # depende de que el LLM lo "recuerde" ni lo repita.
    async with db.session() as s:
        report = (await s.execute(select(Report))).scalars().one()
    content = Path(report.path).read_bytes()
    wb = load_workbook(io.BytesIO(content))
    flat_cells = " ".join(str(c.value) for sheet in wb.worksheets for row in sheet.iter_rows() for c in row if c.value is not None)
    assert "Marta Ruiz" in flat_cells
    assert "1700" in flat_cells or "1.700" in flat_cells.replace("\xa0", " ")
