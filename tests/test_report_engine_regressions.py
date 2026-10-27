"""Pruebas sintéticas del mismo motor de informes; nunca datos reales ni credenciales."""

from __future__ import annotations

# Importar el motor documental sin inicializar toda la aplicación ni sus conectores.
# El paquete empresarial requiere dependencias/servicios externos no disponibles aquí.
import ast
import sys
import types
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest

ROOT = Path(__file__).resolve().parents[1] / "src/mia"
_STUBBED_PACKAGES = ("mia.modules", "mia.modules.rrhh", "mia.modules.rrhh.informes")
# conftest.py (vía mia.main) ya ha importado estos paquetes de verdad antes de que pytest recoja
# este fichero: hay que guardar lo que había (el módulo real, no None) para poder restaurarlo, no
# solo popearlo — popear deja sys.modules sin el paquete real durante el resto de la sesión.
_saved_modules = {fullname: sys.modules.get(fullname) for fullname in _STUBBED_PACKAGES}
for fullname, path in [
    ("mia.modules", ROOT / "modules"),
    ("mia.modules.rrhh", ROOT / "modules/rrhh"),
    ("mia.modules.rrhh.informes", ROOT / "modules/rrhh/informes"),
]:
    module = types.ModuleType(fullname)
    module.__path__ = [str(path)]
    sys.modules[fullname] = module
from mia.modules.rrhh.informes.document import Block, Document, to_document
from mia.modules.rrhh.informes.render import to_pdf, to_pptx, to_xlsx
from mia.modules.rrhh.informes.render_html import to_html

# Deshacer el stub ahora que ya tenemos las referencias que necesitábamos: dejarlo puesto
# envenena sys.modules para el resto de la sesión de pytest (otros test_*.py, recogidos después
# en el mismo proceso, harían `from mia.modules.rrhh import build_rrhh_registry` contra este stub sin
# build_rrhh_registry y fallarían). Los submódulos ya importados (document/render/render_html) quedan
# bien cacheados bajo su propia clave; los 3 paquetes contenedores stubeados se restauran a lo que
# había antes (el paquete real ya cargado por conftest.py), no simplemente se popean.
for _fullname, _previous in _saved_modules.items():
    if _previous is None:
        sys.modules.pop(_fullname, None)
    else:
        sys.modules[_fullname] = _previous


def extract_function(path, name):
    tree = ast.parse((ROOT / path).read_text(encoding="utf-8"))
    function = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == name)
    namespace = {
        "Any": Any,
        "mean": __import__("statistics").mean,
        "median": __import__("statistics").median,
        "LAST_TABLE": "ultima_presentacion",
    }
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(path), "exec"), namespace)
    return namespace[name]


remember_table = extract_function("conversation/tools.py", "remember_table")
_group_stats = extract_function("modules/rrhh/nominas/analisis_economico.py", "_group_stats")
LAST_TABLE = "ultima_presentacion"


NOW = datetime(2026, 10, 14)
SERIES = {
    "type": "line",
    "title": "Comparativa",
    "labels": ["2025", "2026"],
    "series": [{"name": "Sistemas", "values": [123.40, 140.25]}, {"name": "Operaciones", "values": [230.40, 260.25]}],
}


def test_saved_session_rows_not_mutated_and_export_metadata_marks_truncation():
    original = {
        "type": "dashboard",
        "title": "Panel",
        "tables": [
            {"title": "Detalle", "columns": ["n"], "rows": [[n] for n in range(6)]},
        ],
    }
    state = {}
    remember_table(state, original, 3, {"ok": True}, source_tool="some_tool")
    assert len(original["tables"][0]["rows"]) == 6
    assert len(state[LAST_TABLE]["tables"][0]["rows"]) == 3
    assert state[LAST_TABLE]["__export_meta__"]["truncated"]
    assert state[LAST_TABLE]["__export_meta__"]["tables"][0]["total"] == 6


def test_distinct_people_not_person_months():
    stats = _group_stats([100.0] * 6, 2, False, 3)
    assert stats["personas_con_nomina"] == 2
    assert stats["observaciones_nomina"] == 6
    assert _group_stats([1, 2, 3], 1, True, 3)["personas_con_nomina"] == 1


def test_duplicate_graphs_not_repeated():
    d = to_document(
        {
            "type": "table",
            "title": "Datos",
            "columns": ["Depto", "Total"],
            "rows": [["Sistemas", 10], ["Operaciones", 20]],
            "charts": [
                {
                    "type": "donut",
                    "title": "Distribución",
                    "labels": ["A", "B"],
                    "series": [{"name": "Total", "values": [10, 20]}],
                },
                {
                    "type": "bar",
                    "title": "Ranking",
                    "labels": ["A", "B"],
                    "series": [{"name": "Total", "values": [10, 20]}],
                },
            ],
        }
    )
    assert sum(b.kind == "chart" for b in d.blocks) == 1


def test_exports_preserve_chart_series_and_excel_numeric_values():
    doc = Document(
        "Informe sintético",
        [
            Block("chart", "Comparativa", SERIES),
            Block(
                "table",
                "Importes",
                {"columns": ["Nombre", "Importe", "%"], "rows": [["Sistemas", "9.087,57 €", "56,0%"]]},
            ),
            Block("kpis", "Indicadores", {f"Métrica {i}": i for i in range(16)}),
        ],
    )
    xlsx = to_xlsx(doc, NOW)
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(xlsx))
    assert len(wb.worksheets[0]._charts[0].series) == 2
    assert wb.worksheets[1]["B7"].value == 9087.57
    assert wb.worksheets[1]["C7"].value == 0.56
    assert len({ws.title for ws in wb.worksheets}) == len(wb.worksheets)
    pptx = to_pptx(doc, NOW)
    from pptx import Presentation

    ppt = Presentation(BytesIO(pptx))
    charts = [sh.chart for slide in ppt.slides for sh in slide.shapes if sh.has_chart]
    assert charts and len(charts[0].series) == 2
    assert "Métrica 15" in " ".join(sh.text for slide in ppt.slides for sh in slide.shapes if sh.has_text_frame)
    pdf = to_pdf(doc, NOW)
    assert pdf.startswith(b"%PDF")
    html = to_html(doc, NOW)
    assert "Operaciones" in html.decode("utf-8")
    assert "Sistemas" in html.decode("utf-8")


def test_long_pptx_cells_are_not_truncated():
    long = "detalles-importantes-" * 30
    doc = Document("Longitud", [Block("table", "Texto", {"columns": ["Código", "Detalle"], "rows": [["A", long]]})])
    output = to_pptx(doc, NOW)
    from pptx import Presentation

    p = Presentation(BytesIO(output))
    rendered = []
    for slide in p.slides:
        for shape in slide.shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text for cell in row.cells]
                    if cells[0] == "Código":
                        continue
                    rendered.append(cells[1])
    assert "".join(rendered) == long


@pytest.fixture(scope="module")
def _fake_llm_stack():
    """Stub de mia.llm/mia.settings para probar mia.reporting.content_composer sin red ni
    credenciales reales. Restaura sys.modules al salir (teardown): dejar el stub puesto rompería a
    cualquier otro test_*.py recogido después en la misma sesión de pytest que necesite el mia.llm
    real (p. ej. uno que haga ``from mia.llm import ToolCall`` de verdad)."""
    import dataclasses
    import importlib
    import sys
    import types

    saved = {
        name: sys.modules.get(name)
        for name in (
            "mia.llm",
            "mia.llm.base",
            "mia.llm.tiers",
            "mia.settings",
            "mia.reporting.content_composer",
        )
    }
    llm_root = types.ModuleType("mia.llm")
    llm_root.__path__ = [str(ROOT / "llm")]
    sys.modules["mia.llm"] = llm_root
    base = types.ModuleType("mia.llm.base")

    @dataclasses.dataclass
    class Message:
        role: str
        content: str

    @dataclasses.dataclass
    class LLMRequest:
        system: str
        messages: list
        response_schema: dict
        max_output_tokens: int
        model: str | None

    @dataclasses.dataclass
    class TextDelta:
        text: str

    base.LLMProvider, base.Message, base.LLMRequest, base.TextDelta = object, Message, LLMRequest, TextDelta
    sys.modules["mia.llm.base"] = base
    tiers = types.ModuleType("mia.llm.tiers")
    tiers.model_for_tier = lambda _settings, _tier: "gpt-4.1"
    sys.modules["mia.llm.tiers"] = tiers
    settings = types.ModuleType("mia.settings")
    settings.LLMSettings = object
    settings.ModelTier = type("ModelTier", (), {"ADVANCED": "advanced"})
    sys.modules["mia.settings"] = settings
    sys.modules.pop("mia.reporting.content_composer", None)
    composer = importlib.import_module("mia.reporting.content_composer")

    yield composer

    for name, module in saved.items():
        if module is None:
            sys.modules.pop(name, None)
        else:
            sys.modules[name] = module


def test_composer_selects_only_requested_blocks_and_preserves_warnings(_fake_llm_stack):
    # La selección usa el proveedor ya inyectado; se simula su API sin conexión externa.
    import asyncio
    import json

    composer = _fake_llm_stack
    from mia.llm.base import TextDelta

    doc = Document(
        "Informe RRHH",
        [
            Block("table", "Plantilla por departamento", {"columns": ["D", "N"], "rows": [["A", 10]]}),
            Block("table", "Absentismo", {"columns": ["Tipo", "Dias"], "rows": [["IT", 20]]}),
            Block("warnings", "Advertencias", {"items": [{"severity": "info", "code": "X", "message": "Diferencias"}]}),
        ],
    )

    class FakeLLM:
        async def stream(self, request):
            assert request.model == "gpt-4.1"
            assert "20" not in request.messages[0].content  # no importes ni valores a LLM
            yield TextDelta(json.dumps({"selected_indices": [1], "unavailable": []}))

    result = asyncio.run(
        composer.compose_for_request(doc, "Informe solo de absentismo", FakeLLM(), llm_settings=object())
    )
    assert [block.title for block in result.blocks] == ["Absentismo", "Advertencias"]
    assert len(doc.blocks) == 3  # copia inmutable


def test_composer_fails_closed_for_unavailable_requested_metrics(_fake_llm_stack):
    import asyncio
    import json

    composer = _fake_llm_stack
    from mia.llm.base import TextDelta

    class FakeLLM:
        async def stream(self, request):
            yield TextDelta(json.dumps({"selected_indices": [0], "unavailable": ["Coste salarial"]}))

    doc = Document(
        "Ejemplo",
        [
            Block("table", "Empleados", {"columns": ["Nombre"], "rows": [["X"]]}),
            Block("table", "Turnos", {"columns": ["Fecha"], "rows": [["2026"]]}),
        ],
    )
    with pytest.raises(composer.MissingReportEvidence):
        asyncio.run(composer.compose_for_request(doc, "Informe de coste salarial", FakeLLM()))
