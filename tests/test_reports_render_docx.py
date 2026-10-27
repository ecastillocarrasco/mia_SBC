"""Visual-contract tests for the premium Word (DOCX) renderer — same bar as PDF/XLSX/PPTX in
``tests/test_reports_premium_design.py``: marca Sunset Beach Club en cada página, numeración real de
páginas, tarjetas KPI, severidad en color y gráficos incrustados como imagen.

All values are synthetic. No corporate data, network access, or live LLM calls.
"""

from __future__ import annotations

import io
from datetime import datetime

from docx import Document as WordDocument
from docx.shared import RGBColor

from mia.reporting.document import Block, Document
from mia.reporting.render_docx import SEVERITY_COLOR, to_docx

NOW = datetime(2026, 10, 14)


def make_document(categories: int = 3) -> Document:
    labels = [f"Area {i:02}" for i in range(categories)]
    values_a = [1500 + i * 17 for i in range(categories)]
    values_b = [2000 + i * 13 for i in range(categories)]
    return Document(
        "Costes por departamento",
        subtitle="Prueba con cifras sintéticas",
        blocks=[
            Block("kpis", "Indicadores", {"Coste total": "2.000,00 €", "Variación": "5,0 %"}),
            Block(
                "chart",
                "Comparativa",
                {
                    "title": "Comparativa",
                    "type": "grouped_bar",
                    "labels": labels,
                    "series": [{"name": "2025", "values": values_a}, {"name": "2026", "values": values_b}],
                },
            ),
            Block(
                "table",
                "Datos detallados",
                {
                    "columns": ["Area", "2025", "2026"],
                    "rows": [[name, a, b] for name, a, b in zip(labels, values_a, values_b, strict=True)],
                },
            ),
            Block(
                "warnings",
                "Avisos",
                {"items": [{"severity": "blocker", "code": "DEMO", "message": "Cifras sintéticas", "source": "QA"}]},
            ),
        ],
    )


def test_logo_and_brand_rule_repeat_on_every_page_via_the_section_header():
    # En Word, un único encabezado de sección se repite automáticamente en todas las páginas —
    # a diferencia de PDF/PPTX, aquí no hay "páginas" que enumerar en la prueba; se verifica la
    # estructura (una sola sección, sin cabecera de primera página distinta) que garantiza la
    # repetición, igual que ya hacen los tests de PDF ("logo en cada página") y PPTX.
    word = WordDocument(io.BytesIO(to_docx(make_document(10), NOW)))
    assert len(word.sections) == 1
    section = word.sections[0]
    assert section.different_first_page_header_footer is False
    header_xml = section.header._element.xml
    assert "<pic:pic" in header_xml or "blipFill" in header_xml  # el logotipo está incrustado
    assert len(section.header.tables) == 1  # el filete dorado de marca


def test_page_numbers_are_real_fields_not_static_text():
    word = WordDocument(io.BytesIO(to_docx(make_document(3), NOW)))
    footer_xml = word.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml
    assert "fldChar" in footer_xml
    # Word debe recalcular los campos al abrir, no dejar "Página 1 de 1" congelado.
    assert "updateFields" in word.settings.element.xml


def test_kpis_render_as_branded_cards_not_a_generic_table():
    word = WordDocument(io.BytesIO(to_docx(make_document(3), NOW)))
    kpi_table = word.tables[0]
    cell_xml = kpi_table.cell(0, 0)._tc.xml
    assert "D8A565" in cell_xml  # filete superior dorado
    assert "F6F3EE" in cell_xml  # relleno de marca
    values = " ".join(c.text for row in kpi_table.rows for c in row.cells)
    assert "2.000,00" in values and "5,0" in values


def test_large_chart_is_split_into_several_high_resolution_images_not_lost():
    word = WordDocument(io.BytesIO(to_docx(make_document(56), NOW)))
    assert len(word.inline_shapes) == 4  # 56 etiquetas / 15 por imagen, igual que PDF/PPTX


def test_warning_severity_has_a_distinct_color_not_only_the_raw_word():
    word = WordDocument(io.BytesIO(to_docx(make_document(3), NOW)))
    warnings_table = next(t for t in word.tables if "DEMO" in t.cell(0, 0).text)
    run = warnings_table.cell(0, 0).paragraphs[0].runs[0]
    assert run.font.color.rgb == RGBColor.from_string(SEVERITY_COLOR["blocker"])


def test_annex_truncation_is_never_silent():
    doc = Document(
        "Listado",
        blocks=[
            Block(
                "annex",
                "Detalle",
                {"columns": ["Col"], "rows": [[i] for i in range(10)], "total_rows": 5000, "truncated": True},
            ),
        ],
    )
    word = WordDocument(io.BytesIO(to_docx(doc, NOW)))
    text = "\n".join(p.text for p in word.paragraphs)
    assert "Mostrando 10 de 5000 filas" in text


def test_content_width_fits_within_a4_margins():
    from mia.reporting.render_docx import Emu  # noqa: PLC0415

    word = WordDocument(io.BytesIO(to_docx(make_document(3), NOW)))
    section = word.sections[0]
    content_width = Emu(section.page_width - section.left_margin - section.right_margin).inches
    assert 5.5 < content_width < 7.5
