"""Indexación de contenido sensible (nóminas, documentos firmados): mismo patrón que los CVs, por
decisión explícita (ver ``DataSettings.index_sensitive_content``, ``mia.modules.rrhh.document_indexing``).
"""

from __future__ import annotations

import io
from collections.abc import AsyncIterator
from datetime import UTC, datetime
from pathlib import Path

import docx
import httpx
import pytest
from reportlab.pdfgen import canvas
from sqlalchemy import select

from mia.connectors.sharepoint import SharePointClient, StaticTokenProvider
from mia.data import Database, Employee, HrDocumentChunk, HrDocumentText, Payslip, SignedDocument
from mia.llm.fake import FakeLLMProvider
from mia.modules.rrhh.document_indexing import DocumentIndexer
from mia.modules.rrhh.document_search import DocumentSearchIndex
from tests.dbutil import POSTGRES_URL, fresh_database

requires_postgres = pytest.mark.skipif(
    not POSTGRES_URL, reason="pgvector (cosine_distance / DISTINCT ON) solo funciona contra Postgres real"
)


def _pdf(text: str) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    y = 800
    for line in text.splitlines():
        pdf.drawString(40, y, line)
        y -= 16
    pdf.save()
    return buffer.getvalue()


def _scanned_pdf(text: str = "NOMINA DE JULIO CON PLUS DE NOCTURNIDAD INCLUIDO") -> bytes:
    """PDF «escaneado»: el texto vive solo como píxeles de una imagen incrustada (sin capa de
    texto nativa) — desde que el OCR es local (Tesseract, no un LLM al que baste con darle un
    guion), la imagen debe contener texto real y reconocible (>=40 caracteres: el umbral de
    ``extract_document_text_local`` para «ok» frente a «empty») para que el test tenga sentido."""
    import pymupdf
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1300, 310), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 42)
    except OSError:
        font = ImageFont.truetype("DejaVuSans.ttf", 42)
    draw.text((30, 120), text, fill="black", font=font)
    png = io.BytesIO()
    image.save(png, format="PNG")

    doc = pymupdf.open()
    page = doc.new_page(width=1300, height=310)
    page.insert_image(page.rect, stream=png.getvalue())
    data = doc.tobytes()
    doc.close()
    return data


def _docx(text: str) -> bytes:
    document = docx.Document()
    for paragraph in text.split("\n"):
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _files_transport(files: dict[tuple[str, int], bytes]) -> httpx.MockTransport:
    """Simula ``items({id})/File/$value`` para la lista indicada (lo que usa ``download_item_file``)."""

    def handler(request: httpx.Request) -> httpx.Response:
        path = request.url.path
        for (title, item_id), content in files.items():
            if f"getbytitle('{title}')/items({item_id})/File" in path:
                return httpx.Response(200, content=content)
        return httpx.Response(404)

    return httpx.MockTransport(handler)


@pytest.fixture
async def db(tmp_path: Path) -> AsyncIterator[Database]:
    database = await fresh_database(tmp_path, "docidx.db")
    yield database
    await database.dispose()


@requires_postgres
async def test_payslip_pdf_is_extracted_chunked_and_searchable(db: Database) -> None:
    async with db.session() as s:
        s.add(
            Payslip(
                sp_id=501,
                person_key="pk-ana",
                period="2026-07",
                file_name="2026-07_Julio.pdf",
                modified_at=datetime(2026, 7, 1, tzinfo=UTC),
            )
        )
    client = SharePointClient(
        site_url="https://x/sites/rrhh",
        tokens=StaticTokenProvider(),
        transport=_files_transport({("Nominas", 501): _pdf("Camarera de pisos. Plus de nocturnidad incluido.")}),
    )
    indexer = DocumentIndexer(
        domain="nominas",
        list_title="Nominas",
        model=Payslip,
        db=db,
        client=client,
        llm=FakeLLMProvider(),
        embedding_model="x",
    )
    report = await indexer.run()
    await client.aclose()
    assert report.statuses == {"ok": 1}
    async with db.session() as s:
        text = (await s.scalars(select(HrDocumentText))).one()
        chunks = list((await s.scalars(select(HrDocumentChunk))).all())
    assert (text.domain, text.document_id, text.person_key) == ("nominas", 501, "pk-ana")
    assert text.status == "ok" and "nocturnidad" in text.text
    assert len(chunks) == 1 and chunks[0].domain == "nominas"

    index = DocumentSearchIndex("nominas")
    async with db.session() as s:
        vector = (await FakeLLMProvider().embed(["plus de nocturnidad"], "x"))[0]
        hits = await index.search(s, vector, allowed=None, limit=10)
    assert hits and hits[0].document_id == 501


async def test_scanned_payslip_pdf_is_read_visually(db: Database) -> None:
    async with db.session() as s:
        s.add(Payslip(sp_id=502, person_key="pk-ana", modified_at=datetime(2026, 7, 1, tzinfo=UTC)))
    client = SharePointClient(
        site_url="https://x/sites/rrhh",
        tokens=StaticTokenProvider(),
        transport=_files_transport({("Nominas", 502): _scanned_pdf()}),
    )
    indexer = DocumentIndexer(
        domain="nominas",
        list_title="Nominas",
        model=Payslip,
        db=db,
        client=client,
        llm=FakeLLMProvider(),
        embedding_model="x",
    )
    report = await indexer.run()
    await client.aclose()
    assert report.statuses == {"ok": 1}
    async with db.session() as s:
        text = (await s.scalars(select(HrDocumentText))).one()
    assert text.detail is not None and text.detail.startswith("ocr_local:pdf")
    assert "NOCTURNIDAD" in text.text.upper()


async def test_signed_document_docx_is_indexed(db: Database) -> None:
    async with db.session() as s:
        s.add(
            SignedDocument(
                sp_id=601,
                person_key="pk-ana",
                doc_type="Reconocimiento médico",
                file_name="RM_2026.docx",
                modified_at=datetime(2026, 1, 1, tzinfo=UTC),
            )
        )
    client = SharePointClient(
        site_url="https://x/sites/rrhh",
        tokens=StaticTokenProvider(),
        transport=_files_transport(
            {("DocumentosFirmadosV2", 601): _docx("Reconocimiento médico: aptitud laboral sin restricciones.")}
        ),
    )
    indexer = DocumentIndexer(
        domain="documentos_firmados",
        list_title="DocumentosFirmadosV2",
        model=SignedDocument,
        db=db,
        client=client,
        llm=FakeLLMProvider(),
        embedding_model="x",
    )
    report = await indexer.run()
    await client.aclose()
    assert report.statuses == {"ok": 1}
    async with db.session() as s:
        text = (await s.scalars(select(HrDocumentText))).one()
    assert text.domain == "documentos_firmados" and "aptitud laboral" in text.text


async def test_unsupported_file_type_is_recorded_without_chunks(db: Database) -> None:
    async with db.session() as s:
        s.add(
            SignedDocument(
                sp_id=602, person_key="pk-ana", doc_type="Otros", modified_at=datetime(2026, 1, 1, tzinfo=UTC)
            )
        )
    client = SharePointClient(
        site_url="https://x/sites/rrhh",
        tokens=StaticTokenProvider(),
        transport=_files_transport({("DocumentosFirmadosV2", 602): b"not a real document, just garbage bytes"}),
    )
    indexer = DocumentIndexer(
        domain="documentos_firmados",
        list_title="DocumentosFirmadosV2",
        model=SignedDocument,
        db=db,
        client=client,
        llm=FakeLLMProvider(),
        embedding_model="x",
    )
    report = await indexer.run()
    await client.aclose()
    assert report.statuses == {"unsupported": 1}
    async with db.session() as s:
        text = (await s.scalars(select(HrDocumentText))).one()
        chunks = await s.scalar(select(HrDocumentChunk))
    assert text.status == "unsupported" and chunks is None


async def test_changed_modified_at_triggers_reindexing(db: Database) -> None:
    async with db.session() as s:
        s.add(Payslip(sp_id=503, person_key="pk-ana", modified_at=datetime(2026, 7, 1, tzinfo=UTC)))
    files = {("Nominas", 503): _pdf("Texto original de la nomina de julio con plus de transporte.")}
    client = SharePointClient(
        site_url="https://x/sites/rrhh", tokens=StaticTokenProvider(), transport=_files_transport(files)
    )

    def make_indexer() -> DocumentIndexer:
        return DocumentIndexer(
            domain="nominas",
            list_title="Nominas",
            model=Payslip,
            db=db,
            client=client,
            llm=FakeLLMProvider(),
            embedding_model="x",
        )

    first = await make_indexer().run()
    unchanged = await make_indexer().run()
    assert (first.processed, unchanged.processed) == (1, 0)

    # El fichero se reemplaza (nuevo Modified en SharePoint) → debe releerse.
    files[("Nominas", 503)] = _pdf("Texto corregido de la nomina de julio con plus de idiomas.")
    async with db.session() as s:
        slip = await s.get(Payslip, 503)
        slip.modified_at = datetime(2026, 7, 15, tzinfo=UTC)
    changed = await make_indexer().run()
    await client.aclose()
    assert changed.processed == 1 and changed.statuses == {"ok": 1}
    async with db.session() as s:
        text = (await s.scalars(select(HrDocumentText))).one()
    assert "idiomas" in text.text and "transporte" not in text.text


@requires_postgres
async def test_buscar_en_nominas_emits_rag_and_document_diagnostic_events(db: Database) -> None:
    """La herramienta conversacional ``buscar_en_nominas`` debe emitir RAG_SEARCH_START/COMPLETE y
    DOCUMENT_SEARCH/DOCUMENT_RETRIEVAL durante una búsqueda real (no solo en la ruta de error), para que
    un panel de diagnóstico pueda ver qué dominio se consultó, cuántos resultados hubo y con qué confianza."""
    import json
    from types import SimpleNamespace

    from mia.auth import UserContext
    from mia.conversation import ToolContext
    from mia.diagnostics import EventStage
    from mia.diagnostics.store import DiagnosticsStore
    from mia.diagnostics.streaming import LiveBroadcaster
    from mia.diagnostics.trace import begin_trace
    from mia.llm import ToolCall
    from mia.modules.registry import build_registry
    from mia.settings import DiagnosticsSettings

    async with db.session() as s:
        s.add_all(
            [
                Employee(
                    sp_id=801,
                    name="Ana",
                    name_key="ana",
                    active=True,
                    department="Pisos",
                    department_key="pisos",
                    position="Camarera",
                    person_key="pk-ana",
                ),
                Payslip(
                    sp_id=801,
                    person_key="pk-ana",
                    employee_name="Ana",
                    period="2026-07",
                    modified_at=datetime(2026, 7, 1, tzinfo=UTC),
                ),
            ]
        )
    client = SharePointClient(
        site_url="https://x/sites/rrhh",
        tokens=StaticTokenProvider(),
        transport=_files_transport(
            {("Nominas", 801): _pdf("Esta nomina incluye un plus de dietas por desplazamiento al centro.")}
        ),
    )
    await DocumentIndexer(
        domain="nominas",
        list_title="Nominas",
        model=Payslip,
        db=db,
        client=client,
        llm=FakeLLMProvider(),
        embedding_model="x",
    ).run()

    llm = FakeLLMProvider()
    settings = SimpleNamespace(
        data=SimpleNamespace(index_sensitive_content=True, content_enabled=True),
        curriculums=SimpleNamespace(embedding_model="x"),
    )
    registry = build_registry(db, settings=settings, llm=llm, sharepoint=client)
    admin = UserContext("u", "Enrique", roles=frozenset({"mia.admin"}))
    ctx = ToolContext(user=admin, request_id=None, now=None, session_state={})  # type: ignore[arg-type]

    store = DiagnosticsStore(DiagnosticsSettings(enabled=True))
    broadcaster = LiveBroadcaster()
    async with begin_trace(query="prueba", user_id="u", store=store, broadcaster=broadcaster) as trace:
        assert trace is not None
        execution = await registry.execute(
            ToolCall("c1", "buscar_en_nominas", json.dumps({"consulta": "dietas", "nombre": "Ana"})), ctx
        )
    await client.aclose()
    assert execution.result.status.value == "ok"

    saved = store.get(trace.trace_id)
    assert saved is not None
    start_events = [e for e in saved.events if e.stage == EventStage.RAG_SEARCH_START]
    complete_events = [e for e in saved.events if e.stage == EventStage.RAG_SEARCH_COMPLETE]
    search_events = [e for e in saved.events if e.stage == EventStage.DOCUMENT_SEARCH]
    retrieval_events = [e for e in saved.events if e.stage == EventStage.DOCUMENT_RETRIEVAL]
    assert len(start_events) == 1 and start_events[0].data["dominio"] == "nominas"
    assert len(complete_events) == 1
    assert complete_events[0].data["hits"] == 1
    assert complete_events[0].data["top_score"] is not None
    assert len(search_events) == 1 and search_events[0].data["filtro_nombre"] is True
    assert len(retrieval_events) == 1 and retrieval_events[0].data["documentos_recuperados"] == 1


@requires_postgres
async def test_domains_are_isolated_when_sp_id_collides(db: Database) -> None:
    """sp_id NO es único entre bibliotecas: una nómina y un documento firmado pueden compartir el mismo ID."""
    async with db.session() as s:
        s.add_all(
            [
                Payslip(sp_id=700, person_key="pk-ana", modified_at=datetime(2026, 7, 1, tzinfo=UTC)),
                SignedDocument(
                    sp_id=700, person_key="pk-luis", doc_type="Contrato", modified_at=datetime(2026, 7, 1, tzinfo=UTC)
                ),
            ]
        )
    client = SharePointClient(
        site_url="https://x/sites/rrhh",
        tokens=StaticTokenProvider(),
        transport=_files_transport(
            {
                ("Nominas", 700): _pdf("Nomina de Ana con un plus de nocturnidad incluido este mes."),
                ("DocumentosFirmadosV2", 700): _pdf("Contrato de Luis indefinido a jornada completa desde enero."),
            }
        ),
    )
    for domain, list_title, model in (
        ("nominas", "Nominas", Payslip),
        ("documentos_firmados", "DocumentosFirmadosV2", SignedDocument),
    ):
        await DocumentIndexer(
            domain=domain,
            list_title=list_title,
            model=model,
            db=db,
            client=client,
            llm=FakeLLMProvider(),
            embedding_model="x",
        ).run()
    await client.aclose()
    async with db.session() as s:
        texts = {t.domain: t for t in (await s.scalars(select(HrDocumentText))).all()}
    assert len(texts) == 2
    assert "nocturnidad" in texts["nominas"].text and "Luis" in texts["documentos_firmados"].text
    assert texts["nominas"].person_key == "pk-ana" and texts["documentos_firmados"].person_key == "pk-luis"

    nominas_index = DocumentSearchIndex("nominas")
    docs_index = DocumentSearchIndex("documentos_firmados")
    async with db.session() as s:
        vector = (await FakeLLMProvider().embed(["nocturnidad"], "x"))[0]
        nominas_hits = await nominas_index.search(s, vector, allowed=None, limit=10)
        docs_hits = await docs_index.search(s, vector, allowed=None, limit=10)
    assert {h.document_id for h in nominas_hits} == {700}
    assert {h.document_id for h in docs_hits} == {700}  # mismo ID numérico, dominios distintos
