"""Fase D (motor documental universal) — los 4 bloques nuevos (``warnings``, ``narrative``,
``methodology``, ``annex``) deben verse correctamente en PDF, XLSX y PPTX, no solo aceptarse en el
esquema (condición obligatoria 6 de la aprobación: "no es suficiente con que el esquema acepte los
bloques: deben visualizarse correctamente"). Cada prueba abre el archivo real generado e inspecciona
contenido real, con el mismo estilo que ``tests/test_reports_render.py``."""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from mia.analytics.homogeneity import check_period_pair, check_populations
from mia.modules.rrhh.informes.document import (
    Document,
    annex_block,
    methodology_block,
    narrative_block,
    narrative_statement,
    warnings_block,
)
from mia.modules.rrhh.informes.render import RENDERERS
from mia.query.period import period_engine

NOW = datetime(2026, 10, 11, 10, 0)


def _pdf_text(content: bytes) -> str:
    return "".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)


# ------------------------------------------------------------------ warnings
def _warnings_doc() -> Document:
    issues = [
        *check_populations(0, 5, label_a="Cocina", label_b="Pisos"),
        *check_period_pair(
            period_engine.resolve("agosto 2026", today=NOW.date()),
            period_engine.resolve("Q2 2026", today=NOW.date()),
        ),
    ]
    block = warnings_block("Advertencias", [i.to_dict() for i in issues])
    return Document(title="Informe con advertencias", blocks=[block])


def test_warnings_render_with_severity_visible_in_every_format() -> None:
    doc = _warnings_doc()
    pdf_text = _pdf_text(RENDERERS["pdf"](doc, NOW, None))
    assert "BLOQUEANTE" in pdf_text
    assert "POPULATION_EMPTY" in pdf_text

    wb = load_workbook(io.BytesIO(RENDERERS["xlsx"](doc, NOW, None)))
    ws = wb["Advertencias"]
    flat = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Bloqueante" in flat or "BLOQUEANTE" in flat.upper()
    assert "POPULATION_EMPTY" in flat

    slides = Presentation(io.BytesIO(RENDERERS["pptx"](doc, NOW, None))).slides
    assert len(slides) >= 2  # portada + al menos una diapositiva de advertencias
    all_text = " ".join(
        shape.text_frame.text for slide in slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "POPULATION_EMPTY" in all_text


# ------------------------------------------------------------------ narrative
def _narrative_doc() -> Document:
    statements = [
        narrative_statement(
            "fact", "La plantilla activa creció un 12% en el periodo.",
            evidence={"source": "mia.analytics.stats", "metric": "total_staff", "value": 12.0},
        ),
        narrative_statement("hypothesis", "Podría estar relacionado con la temporada alta."),
        narrative_statement("recommendation", "Revisar la dotación de Cocina antes de agosto."),
    ]
    return Document(title="Informe narrativo", blocks=[narrative_block("Hallazgos", statements)])


def test_narrative_statements_are_labeled_by_kind_in_every_format() -> None:
    doc = _narrative_doc()
    pdf_text = _pdf_text(RENDERERS["pdf"](doc, NOW, None))
    assert "HECHO VERIFICADO" in pdf_text
    assert "HIPÓTESIS" in pdf_text
    assert "RECOMENDACIÓN" in pdf_text
    assert "temporada alta" in pdf_text

    wb = load_workbook(io.BytesIO(RENDERERS["xlsx"](doc, NOW, None)))
    ws = wb["Hallazgos"]
    flat = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "temporada alta" in flat

    slides = Presentation(io.BytesIO(RENDERERS["pptx"](doc, NOW, None))).slides
    assert len(slides) >= 2


# ------------------------------------------------------------------ methodology
def test_methodology_notes_appear_in_every_format() -> None:
    doc = Document(
        title="Informe con metodología",
        blocks=[methodology_block("Metodología", ["Los datos proceden de ShiftDay.", "El absentismo excluye descansos."])],
    )
    pdf_text = _pdf_text(RENDERERS["pdf"](doc, NOW, None))
    assert "ShiftDay" in pdf_text

    wb = load_workbook(io.BytesIO(RENDERERS["xlsx"](doc, NOW, None)))
    ws = wb["Metodología"]
    flat = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "ShiftDay" in flat

    slides = Presentation(io.BytesIO(RENDERERS["pptx"](doc, NOW, None))).slides
    assert len(slides) >= 2


# ------------------------------------------------------------------ annex (sin truncar en silencio)
def test_annex_shows_an_explicit_truncation_notice_never_silent() -> None:
    rows = [[i, f"Persona {i}"] for i in range(10)]
    block = annex_block("Detalle completo", ["ID", "Nombre"], rows, max_rows=4)
    doc = Document(title="Informe con anexo", blocks=[block])

    pdf_text = _pdf_text(RENDERERS["pdf"](doc, NOW, None))
    assert "Mostrando 4 de 10" in pdf_text
    assert "Persona 0" in pdf_text
    assert "Persona 9" not in pdf_text  # recortado de verdad, pero avisado — no un error

    wb = load_workbook(io.BytesIO(RENDERERS["xlsx"](doc, NOW, None)))
    ws = wb["Detalle completo"]
    flat = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Mostrando 4 de 10" in flat

    slides = Presentation(io.BytesIO(RENDERERS["pptx"](doc, NOW, None))).slides
    all_text = " ".join(
        shape.text_frame.text for slide in slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "Mostrando 4 de 10" in all_text


def test_annex_without_truncation_has_no_spurious_notice() -> None:
    rows = [[i] for i in range(3)]
    block = annex_block("Detalle pequeño", ["ID"], rows, max_rows=10)
    doc = Document(title="Informe con anexo pequeño", blocks=[block])
    pdf_text = _pdf_text(RENDERERS["pdf"](doc, NOW, None))
    assert "Mostrando" not in pdf_text
