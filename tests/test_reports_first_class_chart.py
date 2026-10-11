"""Bloque de gráfico de primera clase (Bloque II, sección 10, condición obligatoria 3 de la aprobación
de Fase A): ``Block(kind="chart", ...)`` construido explícitamente por el backend (con
``mia.modules.rrhh.charts``), no derivado heurísticamente por cada renderizador por separado.

Antes de este cambio, SOLO existían bloques ``table`` de 2 columnas de los que el renderizador
"adivinaba" un gráfico en el momento de pintar (``chart_spec()`` en cada uno de los tres
renderizadores). Eso sigue funcionando igual (ver ``test_reports_render.py``) — este archivo prueba el
camino NUEVO: un gráfico que el backend ya decidió explícitamente, reutilizado sin volver a derivarlo,
en los tres formatos."""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from mia.modules.rrhh.charts import bar_chart, donut_chart
from mia.modules.rrhh.informes.document import Block, Document
from mia.modules.rrhh.informes.render import RENDERERS, _chart_of, _rows

NOW = datetime(2026, 10, 11, 10, 0)


def _doc_with_chart() -> Document:
    spec = bar_chart("Plantilla por departamento", ["Cocina", "Pisos", "Spa"], [12.0, 8.0, 3.0])
    return Document(title="Informe de prueba", blocks=[Block("chart", "Plantilla por departamento", spec)])


def test_chart_of_returns_the_explicit_spec_without_re_deriving_it() -> None:
    """A diferencia de un bloque ``table``, el gráfico no se vuelve a calcular: es exactamente el dict
    que construyó el backend."""
    spec = donut_chart("Composición", ["A", "B"], [70.0, 30.0])
    block = Block("chart", "Composición", spec)
    assert _chart_of(block) is spec  # el mismo objeto, no una copia recalculada


def test_rows_reconstructs_a_consistent_data_table_from_the_chart_spec() -> None:
    spec = bar_chart("Plantilla por departamento", ["Cocina", "Pisos"], [12.0, 8.0], series_name="Personas")
    columns, rows = _rows(Block("chart", "Plantilla por departamento", spec))
    assert columns == ["Categoría", "Personas"]
    assert rows == [["Cocina", 12.0], ["Pisos", 8.0]]


def test_chart_block_renders_in_every_format_with_the_same_source_data() -> None:
    doc = _doc_with_chart()
    for fmt in ("pdf", "xlsx", "pptx"):
        content = RENDERERS[fmt](doc, NOW, None)
        assert content  # se generó sin lanzar excepción, en los tres formatos

    pdf_text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(RENDERERS["pdf"](doc, NOW, None))).pages)
    assert "Plantilla por departamento" in pdf_text
    assert "Cocina" in pdf_text  # la tabla reconstruida bajo el gráfico también se pinta

    wb = load_workbook(io.BytesIO(RENDERERS["xlsx"](doc, NOW, None)))
    assert wb.sheetnames == ["Plantilla por departamento"]
    ws = wb["Plantilla por departamento"]
    flat = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Cocina" in flat and "Categoría" in flat

    slides = Presentation(io.BytesIO(RENDERERS["pptx"](doc, NOW, None))).slides
    assert len(slides) >= 3  # portada + diapositiva de gráfico + diapositiva(s) de tabla
