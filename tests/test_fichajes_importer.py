"""Importador del Excel de Conectia: cruce de persona por nombre (sin depender del orden ni de
tildes/ñ), idempotencia al reimportar, y filas sin persona reconocida reportadas, nunca
descartadas en silencio. Las cabeceras y el formato de nombre replican los ficheros reales
(semana 41, PISOS/RESTAURANTE)."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from mia.data import ClockMark
from mia.modules.rrhh.fichajes.importer import import_directory, import_from_attachment
from tests.dbutil import fresh_database
from tests.hr_seed import _emp

_HEADER = (
    "Fecha", "Área", "Código del empleado", "Nombre completo", "Hora marcaje",
    "Origen del marcaje", "Motivo", "Nombre Terminal", "Tipo marcaje", "Link localización",
)


def _write_xlsx(path: Path, rows: list[tuple]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(_HEADER)
    for row in rows:
        ws.append(row)
    wb.save(path)


def _row(fecha: datetime, codigo: str, nombre: str, *, origen="Normal", terminal="Terminal -2", tipo="App") -> tuple:
    return (fecha, "1-1", codigo, nombre, fecha.strftime("%H:%M"), origen, "", terminal, tipo, "")


def _xlsx_bytes(rows: list[tuple]) -> bytes:
    """Mismo contenido que ``_write_xlsx``, pero en memoria — simula el adjunto subido en el chat,
    sin tocar el disco (``import_from_attachment`` nunca lee de un directorio)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(_HEADER)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_matches_by_name_regardless_of_word_order(tmp_path: Path) -> None:
    db = await fresh_database(tmp_path, "fichajes_import_1.db")
    async with db.session() as s:
        s.add(_emp(1, "Melania Rosello Diaz", "Restaurante", person_key="pk-melania"))
    excel_dir = tmp_path / "excels"
    excel_dir.mkdir()
    _write_xlsx(excel_dir / "semana.xlsx", [_row(datetime(2026, 10, 5, 9, 19), "1-001080", "ROSELLO DIAZ, MELANIA")])
    report = await import_directory(db, excel_dir, now=datetime.now(UTC))
    assert report.imported == 1
    assert report.unmatched == {}
    assert report.ambiguous == {}
    async with db.session() as s:
        mark = (await s.scalars(select(ClockMark))).one()
        assert mark.person_key == "pk-melania"
    await db.dispose()


async def test_matches_by_name_without_accents_or_enie(tmp_path: Path) -> None:
    """Aviso explícito del usuario: en turnos (y por extensión, cualquier fuente) los nombres
    pueden venir sin tildes ni ñ. normalize_key() ya lo resuelve (NFKD + descartar combining)."""
    db = await fresh_database(tmp_path, "fichajes_import_2.db")
    async with db.session() as s:
        s.add(_emp(1, "Juan Nuñez Pérez", "Pisos", person_key="pk-nunez"))
    excel_dir = tmp_path / "excels"
    excel_dir.mkdir()
    # Conectia sin tildes ni eñe, como advirtió el usuario.
    _write_xlsx(excel_dir / "semana.xlsx", [_row(datetime(2026, 10, 5, 9, 0), "1-000001", "NUNEZ PEREZ, JUAN")])
    report = await import_directory(db, excel_dir, now=datetime.now(UTC))
    assert report.imported == 1
    assert report.unmatched == {}
    await db.dispose()


async def test_ambiguous_surname_fragment_does_not_cause_a_wrong_match(tmp_path: Path) -> None:
    """"de la rosa" aparece en dos empleados reales (Conectia, semana 41) — el cruce por nombre
    COMPLETO (conjunto de palabras) debe distinguirlos sin ambigüedad."""
    db = await fresh_database(tmp_path, "fichajes_import_3.db")
    async with db.session() as s:
        s.add(_emp(1, "Francisco Javier Cordero de la Rosa", "Pisos", person_key="pk-francisco"))
        s.add(_emp(2, "Patricia de la Rosa Sanchez", "Restaurante", person_key="pk-patricia"))
    excel_dir = tmp_path / "excels"
    excel_dir.mkdir()
    _write_xlsx(excel_dir / "semana.xlsx", [
        _row(datetime(2026, 10, 6, 15, 56), "1-001105", "DE LA ROSA SANCHEZ, PATRICIA"),
    ])
    report = await import_directory(db, excel_dir, now=datetime.now(UTC))
    assert report.imported == 1
    assert report.unmatched == {}
    assert report.ambiguous == {}
    async with db.session() as s:
        mark = (await s.scalars(select(ClockMark))).one()
        assert mark.person_key == "pk-patricia"
    await db.dispose()


async def test_matches_when_conectia_has_an_extra_middle_name_mia_does_not(tmp_path: Path) -> None:
    """Caso real (importación contra datos reales, 2026-10-19): mIA guarda "Melania Rosello Diaz"
    pero Conectia exporta "MELANIA JACINTA ROSELLO DIAZ" (con segundo nombre). El nombre de mIA es
    un SUBCONJUNTO del de Conectia y es el único candidato así: debe cruzar."""
    db = await fresh_database(tmp_path, "fichajes_import_7.db")
    async with db.session() as s:
        s.add(_emp(1, "Melania Rosello Diaz", "Restaurante", person_key="pk-melania"))
    excel_dir = tmp_path / "excels"
    excel_dir.mkdir()
    _write_xlsx(excel_dir / "semana.xlsx", [
        _row(datetime(2026, 10, 5, 9, 19), "1-001080", "ROSELLO DIAZ, MELANIA JACINTA"),
    ])
    report = await import_directory(db, excel_dir, now=datetime.now(UTC))
    assert report.imported == 1
    assert report.unmatched == {}
    async with db.session() as s:
        mark = (await s.scalars(select(ClockMark))).one()
        assert mark.person_key == "pk-melania"
    await db.dispose()


async def test_subset_match_ambiguous_between_two_employees_is_reported_not_guessed(tmp_path: Path) -> None:
    """Si DOS empleados distintos son cada uno un subconjunto válido del nombre de Conectia, no se
    adivina cuál es: se reporta como ambiguo."""
    db = await fresh_database(tmp_path, "fichajes_import_8.db")
    async with db.session() as s:
        s.add(_emp(1, "Maria Garcia Lopez", "Pisos", person_key="pk-maria"))
        s.add(_emp(2, "Jose Garcia Lopez", "Restaurante", person_key="pk-jose"))
    excel_dir = tmp_path / "excels"
    excel_dir.mkdir()
    _write_xlsx(excel_dir / "semana.xlsx", [
        _row(datetime(2026, 10, 5, 9, 0), "1-000002", "GARCIA LOPEZ, MARIA JOSE"),
    ])
    report = await import_directory(db, excel_dir, now=datetime.now(UTC))
    assert report.imported == 0
    assert report.ambiguous == {"GARCIA LOPEZ, MARIA JOSE": 1}
    await db.dispose()


async def test_known_spelling_variants_are_corrected(tmp_path: Path) -> None:
    """Casos reales confirmados por el usuario (importación semana 41): Jenifer/Jennifer (x2
    apellidos distintos), Vanesa/Vanessa, "de"/"da" Conceicao — siempre por nombre completo."""
    db = await fresh_database(tmp_path, "fichajes_import_9.db")
    async with db.session() as s:
        s.add(_emp(1, "Jennifer Jimenez Jimenez", "Pisos", person_key="pk-jennifer-jj"))
        s.add(_emp(2, "Vanessa Sanchez Valle", "Restaurante", person_key="pk-vanessa"))
        s.add(_emp(3, "Cristiane da Conceicao de Lima", "Pisos", person_key="pk-cristiane"))
    excel_dir = tmp_path / "excels"
    excel_dir.mkdir()
    _write_xlsx(excel_dir / "semana.xlsx", [
        _row(datetime(2026, 10, 5, 9, 0), "1-1", "JIMENEZ JIMENEZ, JENIFER"),
        _row(datetime(2026, 10, 5, 9, 0), "1-2", "SANCHEZ VALLE, VANESA"),
        _row(datetime(2026, 10, 5, 9, 0), "1-3", "DE CONCEICAO DE LIMA, CRISTIANE"),
    ])
    report = await import_directory(db, excel_dir, now=datetime.now(UTC))
    assert report.imported == 3
    assert report.unmatched == {}
    async with db.session() as s:
        keys = {m.person_key for m in (await s.scalars(select(ClockMark))).all()}
        assert keys == {"pk-jennifer-jj", "pk-vanessa", "pk-cristiane"}
    await db.dispose()


async def test_spelling_correction_is_scoped_to_the_exact_full_name_not_the_word(tmp_path: Path) -> None:
    """Incidencia real encontrada al probar contra datos reales: corregir "vanesa"->"vanessa"
    como palabra suelta rompía a "Yesica Vanesa Jofre", una persona DISTINTA que legítimamente se
    llama Vanesa (con una sola ese) — ambas deben seguir cruzando correctamente a la vez."""
    db = await fresh_database(tmp_path, "fichajes_import_11.db")
    async with db.session() as s:
        s.add(_emp(1, "Vanessa Sanchez Valle", "Restaurante", person_key="pk-vanessa"))
        s.add(_emp(2, "Yesica Vanesa Jofre", "Pisos", person_key="pk-yesica"))
    excel_dir = tmp_path / "excels"
    excel_dir.mkdir()
    _write_xlsx(excel_dir / "semana.xlsx", [
        _row(datetime(2026, 10, 5, 9, 0), "1-2", "SANCHEZ VALLE, VANESA"),  # necesita corrección
        _row(datetime(2026, 10, 7, 8, 0), "1-9", "JOFRE AGULLO, YESICA VANESA"),  # NO debe corregirse
    ])
    report = await import_directory(db, excel_dir, now=datetime.now(UTC))
    assert report.imported == 2
    assert report.unmatched == {}
    async with db.session() as s:
        keys = {m.person_key for m in (await s.scalars(select(ClockMark))).all()}
        assert keys == {"pk-vanessa", "pk-yesica"}
    await db.dispose()


async def test_de_da_correction_does_not_affect_unrelated_de_la_rosa_names(tmp_path: Path) -> None:
    """La corrección de Cristiane es del NOMBRE COMPLETO, nunca una regla "de"->"da" genérica:
    no debe tocar a nadie más que use "de" legítimamente."""
    db = await fresh_database(tmp_path, "fichajes_import_10.db")
    async with db.session() as s:
        s.add(_emp(1, "Patricia de la Rosa Sanchez", "Restaurante", person_key="pk-patricia"))
    excel_dir = tmp_path / "excels"
    excel_dir.mkdir()
    _write_xlsx(excel_dir / "semana.xlsx", [
        _row(datetime(2026, 10, 6, 15, 56), "1-001105", "DE LA ROSA SANCHEZ, PATRICIA"),
    ])
    report = await import_directory(db, excel_dir, now=datetime.now(UTC))
    assert report.imported == 1
    assert report.unmatched == {}
    await db.dispose()


async def test_unmatched_name_is_reported_not_silently_dropped(tmp_path: Path) -> None:
    db = await fresh_database(tmp_path, "fichajes_import_4.db")
    async with db.session() as s:
        s.add(_emp(1, "Melania Rosello Diaz", "Restaurante", person_key="pk-melania"))
    excel_dir = tmp_path / "excels"
    excel_dir.mkdir()
    _write_xlsx(excel_dir / "semana.xlsx", [
        _row(datetime(2026, 10, 5, 9, 0), "1-999999", "PERSONA INEXISTENTE, NADIE"),
    ])
    report = await import_directory(db, excel_dir, now=datetime.now(UTC))
    assert report.imported == 0
    assert report.unmatched == {"PERSONA INEXISTENTE, NADIE": 1}
    await db.dispose()


async def test_reimporting_the_same_file_does_not_duplicate(tmp_path: Path) -> None:
    db = await fresh_database(tmp_path, "fichajes_import_5.db")
    async with db.session() as s:
        s.add(_emp(1, "Melania Rosello Diaz", "Restaurante", person_key="pk-melania"))
    excel_dir = tmp_path / "excels"
    excel_dir.mkdir()
    _write_xlsx(excel_dir / "semana.xlsx", [_row(datetime(2026, 10, 5, 9, 19), "1-001080", "ROSELLO DIAZ, MELANIA")])
    r1 = await import_directory(db, excel_dir, now=datetime.now(UTC))
    r2 = await import_directory(db, excel_dir, now=datetime.now(UTC))
    assert r1.imported == 1
    assert r2.imported == 1  # se vuelve a "importar" (merge), pero no duplica la fila
    async with db.session() as s:
        total = (await s.execute(select(func.count()).select_from(ClockMark))).scalar()
        assert total == 1
    await db.dispose()


async def test_import_from_attachment_matches_the_same_way_as_the_directory_path(tmp_path: Path) -> None:
    """``import_from_attachment`` (el camino real desde un adjunto subido en el chat) usa el mismo
    cruce por nombre y el mismo upsert que ``import_directory`` (pruebas) — nunca un segundo motor."""
    db = await fresh_database(tmp_path, "fichajes_import_attachment_1.db")
    async with db.session() as s:
        s.add(_emp(1, "Melania Rosello Diaz", "Restaurante", person_key="pk-melania"))
    content = _xlsx_bytes([_row(datetime(2026, 10, 5, 9, 19), "1-001080", "ROSELLO DIAZ, MELANIA")])
    report = await import_from_attachment(db, content, "fichajes_semana41.xlsx", now=datetime.now(UTC))
    assert report.imported == 1
    assert report.files == 1
    assert report.unmatched == {}
    async with db.session() as s:
        mark = (await s.scalars(select(ClockMark))).one()
        assert mark.person_key == "pk-melania"
    await db.dispose()


async def test_import_from_attachment_reports_unmatched_rows_too(tmp_path: Path) -> None:
    db = await fresh_database(tmp_path, "fichajes_import_attachment_2.db")
    async with db.session() as s:
        s.add(_emp(1, "Melania Rosello Diaz", "Restaurante", person_key="pk-melania"))
    content = _xlsx_bytes([_row(datetime(2026, 10, 5, 9, 0), "1-999999", "PERSONA INEXISTENTE, NADIE")])
    report = await import_from_attachment(db, content, "fichajes.xlsx", now=datetime.now(UTC))
    assert report.imported == 0
    assert report.unmatched == {"PERSONA INEXISTENTE, NADIE": 1}
    await db.dispose()


async def test_missing_expected_column_raises_a_clear_error_not_a_silent_skip(tmp_path: Path) -> None:
    db = await fresh_database(tmp_path, "fichajes_import_6.db")
    excel_dir = tmp_path / "excels"
    excel_dir.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(("Fecha", "Nombre completo"))  # faltan columnas esperadas
    ws.append((datetime(2026, 10, 5, 9, 0), "ALGUIEN, NOMBRE"))
    wb.save(excel_dir / "semana.xlsx")
    with pytest.raises(ValueError, match="faltan columnas"):
        await import_directory(db, excel_dir, now=datetime.now(UTC))
    await db.dispose()
