"""Visual-contract tests for the unified Sunset Beach Club report engine.

All values are synthetic. No corporate data, network access, or live LLM calls.
"""

from __future__ import annotations

import sys
import types
from datetime import datetime
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1] / "src/mia"
# conftest.py (vía mia.main) ya ha importado estos paquetes de verdad antes de que pytest recoja
# este fichero: hay que guardar lo que había (el módulo real, no None) para poder restaurarlo. El
# guard "if fullname not in sys.modules" de antes asumía que, si ya estaba presente, era el real —
# falso cuando otro fichero de test (test_report_engine_regressions.py) lo había dejado stubeado o
# ausente; eso dejaba el stub puesto para el resto de la sesión de pytest sin limpiarlo nunca.
_STUBBED_PACKAGES = ("mia.modules", "mia.modules.rrhh", "mia.modules.rrhh.informes")
_saved_modules = {fullname: sys.modules.get(fullname) for fullname in _STUBBED_PACKAGES}
for fullname, path in (
    ("mia.modules", ROOT / "modules"),
    ("mia.modules.rrhh", ROOT / "modules/rrhh"),
    ("mia.modules.rrhh.informes", ROOT / "modules/rrhh/informes"),
):
    module = types.ModuleType(fullname)
    module.__path__ = [str(path)]
    sys.modules[fullname] = module

from mia.modules.rrhh.informes.document import Block, Document
from mia.modules.rrhh.informes.render import _chart_segments, to_pdf, to_pptx, to_xlsx
from mia.modules.rrhh.informes.render_html import to_html

for _fullname, _previous in _saved_modules.items():
    if _previous is None:
        sys.modules.pop(_fullname, None)
    else:
        sys.modules[_fullname] = _previous

NOW = datetime(2026, 10, 14)


def make_document(categories: int = 3) -> Document:
    labels = [f"Area {i:02}" for i in range(categories)]
    values_a = [1500 + i * 17 for i in range(categories)]
    values_b = [2000 + i * 13 for i in range(categories)]
    return Document(
        "Costes por departamento",
        subtitle="Prueba con cifras sintéticas",
        blocks=[
            Block("kpis", "Indicadores", {"Coste total": "2.000,00 €", "Variación": "5,0 %"}),
            Block(
                "chart",
                "Comparativa",
                {
                    "title": "Comparativa",
                    "type": "grouped_bar",
                    "labels": labels,
                    "series": [{"name": "2025", "values": values_a}, {"name": "2026", "values": values_b}],
                },
            ),
            Block(
                "table",
                "Datos detallados",
                {
                    "columns": ["Area", "2025", "2026"],
                    "rows": [[name, a, b] for name, a, b in zip(labels, values_a, values_b)],
                },
            ),
            Block(
                "warnings",
                "Avisos",
                {"items": [{"severity": "info", "code": "DEMO", "message": "Cifras sintéticas", "source": "QA"}]},
            ),
        ],
    )


def test_chart_segments_keep_all_labels_values_and_series():
    source = make_document(56).blocks[1].data
    chunks = _chart_segments(source, 15)
    assert len(chunks) == 4
    assert [label for chunk in chunks for label in chunk["labels"]] == source["labels"]
    for series_index in range(2):
        assert [v for chunk in chunks for v in chunk["series"][series_index]["values"]] == source["series"][
            series_index
        ]["values"]


def test_pptx_logo_on_every_slide_and_chart_series_preserved():
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    output = to_pptx(make_document(56), NOW)
    deck = Presentation(BytesIO(output))
    assert len(deck.slides) >= 8
    assert all(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes) for slide in deck.slides)
    charts = [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 5  # 56 categories / 13 per slide
    assert all(len(chart.series) == 2 for chart in charts)
    assert sum(len(chart.plots[0].categories) for chart in charts) == 56


def test_xlsx_kpis_link_to_real_cells_and_data_remain_numeric():
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(to_xlsx(make_document(10), NOW)))
    first = workbook.worksheets[0]
    assert first["E8"].value == "=B7"
    assert isinstance(first["B7"].value, (float, int))
    assert first.sheet_view.showGridLines is False
    assert first["C2"].value == "Costes por departamento"
    data_sheet = next(ws for ws in workbook.worksheets if ws["A6"].value == "Area")
    assert len(data_sheet._charts) == 0
    assert data_sheet["B7"].value == 1500
    chart_sheet = next(ws for ws in workbook.worksheets if len(ws._charts) == 1)
    assert len(chart_sheet._charts[0].series) == 2


def test_pdf_has_logo_on_every_page_and_long_table_is_not_lost():
    import fitz

    stream = to_pdf(make_document(56), NOW)
    pages = fitz.open(stream=stream, filetype="pdf")
    assert len(pages) >= 3
    assert all(len(p.get_images()) > 0 for p in pages)
    all_text = "\n".join(page.get_text() for page in pages)
    assert "Area 00" in all_text and "Area 55" in all_text


def test_html_escaped_mixed_signed_bars_and_mandatory_embedded_logo():
    doc = make_document()
    doc.blocks[1].data["series"][0]["values"] = [10, -15, 20]
    doc.title = "<script>alert(1)</script>"
    html = to_html(doc, NOW).decode("utf-8")
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html
    assert "data:image/png;base64," in html
    assert "zero" not in html  # rendered geometry rather than diagnostics
    assert 'x1="330"' in html  # signed-axis origin, not absolute-value reversal
